"""Turn stored results into tables, and export tables / results to Excel."""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

import numpy as np

from spectro.core.validation import describe


def _num(v: Any) -> Any:
    if isinstance(v, (np.floating, np.integer)):
        return v.item()
    return v


def result_table(project, data: dict) -> tuple[list[str], list[list[Any]], list[str]]:
    """(headers, rows, summary lines) for a stored result.

    Concentration results (``ids`` + ``found``) become found / taken /
    recovery columns per compound; anything else is flattened to key/value."""
    ids, found = data.get("ids"), data.get("found")
    if ids and isinstance(found, list) and found and isinstance(found[0], dict):
        comps = [data.get("X", "X"), data.get("Y", "Y")]
        found = [[f.get("X"), f.get("Y")] for f in found]
    else:
        comps = data.get("compounds") or ([data["method"]["compound"]]
                                          if isinstance(data.get("method"), dict)
                                          and data["method"].get("compound") else None)
    if ids and comps and isinstance(found, list) and found and isinstance(found[0], list):
        headers = ["Spectrum"] + [h for c in comps for h in
                                  (f"{c} found", f"{c} taken", f"{c} recovery %")]
        excluded = {int(i) for i in data.get("excluded_ids") or []}
        rows, recs = [], {c: [] for c in comps}
        for sid, f in zip(ids, found):
            try:
                rec = project.record(int(sid))
            except (KeyError, TypeError, ValueError):
                continue
            out = int(sid) in excluded
            row: list[Any] = [rec.name + (" (excluded)" if out else "")]
            for c, v in zip(comps, f):
                t = rec.concentrations.get(c)
                r = 100 * v / t if (t and v is not None) else None
                if r is not None and not out:
                    recs[c].append(r)
                row += [_num(v), t, r]
            rows.append(row)
        summary = []
        for c, v in recs.items():
            if len(v) > 1:
                d = describe(v)
                summary.append(f"{c}: mean recovery {d['mean']:.2f} %, SD {d['sd']:.3f}, "
                               f"RSD {d['rsd']:.2f} % (n = {d['n']})")
        if excluded:
            n_ex = len(excluded & {int(i) for i in ids})
            summary.append(f"{n_ex} {'spectrum' if n_ex == 1 else 'spectra'} excluded from "
                           "the statistics" + (f": {data['exclusion_note']}"
                                               if data.get("exclusion_note") else ""))
        if data.get("notes"):
            summary.append(f"Notes: {data['notes']}")
        return headers, rows, summary
    rows = [[k, v] for k, v in _flatten(data)]
    return ["Item", "Value"], rows, []


def _flatten(obj: Any, prefix: str = "") -> list[tuple[str, Any]]:
    """Nested dict/list → (dotted key, scalar) pairs; long raw arrays skipped."""
    if isinstance(obj, dict):
        out: list[tuple[str, Any]] = []
        for k, v in obj.items():
            if k not in ("x", "y", "residuals", "steps"):
                out += _flatten(v, f"{prefix}.{k}" if prefix else str(k))
        return out
    if isinstance(obj, list):
        if all(not isinstance(v, (dict, list)) for v in obj):
            return [(prefix, ", ".join(_fmt(v) for v in obj[:30])
                     + (" …" if len(obj) > 30 else ""))]
        return [pair for i, v in enumerate(obj[:200]) for pair in _flatten(v, f"{prefix}[{i}]")]
    return [(prefix, _num(obj))]


def _fmt(v: Any) -> str:
    if isinstance(v, float):
        return f"{v:.6g}"
    return str(v)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:28] or "Sheet"
    n, cand = 1, base
    while cand.lower() in used:
        n += 1
        cand = f"{base[:25]} ({n})"
    used.add(cand.lower())
    return cand


def _write_table(ws, headers: list[str], rows: list[list[Any]], start_row: int = 1) -> int:
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="F0EFEC")
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font, c.fill = bold, fill
    for i, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            v = _num(v)
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                v = None
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                c.number_format = "0.0000"
    for j, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(_fmt(r[j - 1])) for r in rows[:200] if j - 1 < len(r)])
        ws.column_dimensions[ws.cell(row=start_row, column=j).column_letter].width = min(60, width + 2)
    return start_row + len(rows) + 1


def export_table_excel(path: str | Path, headers: list[str], rows: list[list[Any]],
                       title: str = "Table", notes: list[str] | None = None) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _sheet_name(title, set())
    end = _write_table(ws, headers, rows)
    for i, line in enumerate(notes or []):
        ws.cell(row=end + 1 + i, column=1, value=line)
    wb.save(str(path))


def export_results_excel(project, path: str | Path, trial_id: int | None = None,
                         result_ids: list[int] | None = None) -> int:
    """One workbook: an index sheet, a methods sheet and one sheet per result."""
    import openpyxl
    from openpyxl.styles import Font

    results = [r for r in project.results(trial_id)
               if result_ids is None or r["id"] in result_ids]
    wb = openpyxl.Workbook()
    idx = wb.active
    idx.title = "Index"
    idx["A1"] = project.meta("name")
    idx["A1"].font = Font(bold=True, size=13)
    idx["A2"] = f"Exported by {project.user}@{project.host}"
    used = {"index", "methods"}
    rows = []
    for r in results:
        sheet = _sheet_name(f"{r['id']} {r['name']}", used)
        headers, trows, summary = result_table(project, r["data"])
        ws = wb.create_sheet(sheet)
        ws["A1"] = r["name"]
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = f"{r['kind']} · result #{r['id']} · {r['created_utc'][:19]} UTC"
        end = _write_table(ws, headers, trows, start_row=4)
        for i, line in enumerate(summary):
            ws.cell(row=end + 1 + i, column=1, value=line)
        rows.append([r["id"], r["name"], r["kind"], r["method_id"], r["created_utc"][:19], sheet])
    _write_table(idx, ["#", "Result", "Kind", "Method #", "Created (UTC)", "Sheet"], rows, 4)
    ms = wb.create_sheet("Methods")
    mrows = []
    for m in project.methods(trial_id):
        d = m["definition"]
        reg = d.get("regression") or {}
        mrows.append([m["id"], m["name"], m["type"], d.get("compound") or ", ".join(d.get("compounds", [])),
                      reg.get("slope"), reg.get("intercept"), reg.get("r"), reg.get("lod"),
                      reg.get("loq")])
    _write_table(ms, ["#", "Method", "Type", "Compound(s)", "Slope", "Intercept", "r", "LOD", "LOQ"],
                 mrows)
    wb.save(str(path))
    return len(results)

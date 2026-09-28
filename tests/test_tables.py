import numpy as np
import pytest
import openpyxl

from spectro.demo import build_demo_project
from spectro.storage.tables import export_results_excel, export_table_excel, result_table


def test_result_tables_and_excel(tmp_path):
    p = build_demo_project(tmp_path / "d.spectro")
    try:
        res = p.results()
        headers, rows, summary = result_table(p, res[-1]["data"])
        assert headers[:4] == ["Spectrum", "PAR found", "PAR taken", "PAR recovery %"]
        assert len(rows) == 8 and summary
        h2, r2, _ = result_table(p, {"t": {"t": 1.2, "p": 0.3}, "values": [1, 2]})
        assert h2 == ["Item", "Value"] and ["t.t", 1.2] in r2
        out = tmp_path / "results.xlsx"
        assert export_results_excel(p, out) == 4
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames[0] == "Index" and "Methods" in wb.sheetnames
        assert len(wb.sheetnames) == 6
        export_table_excel(tmp_path / "t.xlsx", ["a", "b"], [[1, 2.5]], "My: table")
        assert openpyxl.load_workbook(tmp_path / "t.xlsx").active["B2"].value == 2.5
    finally:
        p.close()


def test_model_file_roundtrip(tmp_path):
    import json

    from spectro.core.multicomponent import SpectralModel
    from spectro.core.univariate import UnivariateMethod
    from spectro.storage.modelfile import export_model, import_model
    from spectro.storage.project import Project

    src = build_demo_project(tmp_path / "a.spectro")
    dst = Project.create(tmp_path / "b.spectro")
    try:
        methods = {m["name"]: m for m in src.methods()}
        rd = methods["Ratio difference — PAR"]
        pls = next(m for n, m in methods.items() if n.startswith("PLS2"))
        tabs = [r.id for r in src.records() if "Panadol" in r.name]
        exc = [r.id for r in src.records() if "Excedrin" in r.name]
        for md, ids in ((rd, tabs), (pls, exc)):
            doc = json.loads(json.dumps(export_model(src, md)))
            mid = import_model(dst, doc, None, "x.spmodel")
            d = next(m for m in dst.methods() if m["id"] == mid)["definition"]
            spectra = src.spectra(ids)
            if d["type"] == "univariate":
                assert len(doc["embedded_spectra"]) == 1
                m_src = UnivariateMethod.from_dict(md["definition"])
                m_dst = UnivariateMethod.from_dict(d)
                a = [m_src.predict(s, src.resolver()) for s in spectra]
                b = [m_dst.predict(s, dst.resolver()) for s in spectra]
            else:
                a = SpectralModel.from_dict(md["definition"]).predict(spectra)
                b = SpectralModel.from_dict(d).predict(spectra)
            assert np.allclose(a, b)
        doc["sha256"] = "0" * 64
        with pytest.raises(ValueError):
            import_model(dst, doc, None)
        assert dst.verify()["ok"]
    finally:
        src.close()
        dst.close()

"""Method families, the method guide and the screenshot tools stay consistent
with the app."""

import subprocess
import sys
from pathlib import Path

import pytest

from spectro.core import univariate as uv
from spectro.core.catalog import CATEGORIES, DESCRIPTIONS
from spectro.core.operations import REGISTRY

ROOT = Path(__file__).resolve().parents[1]
GUIDE = ROOT / "docs" / "methods" / "README.md"


def test_every_template_belongs_to_a_family():
    assert set(DESCRIPTIONS) == set(CATEGORIES)
    for name, t in uv.TEMPLATES.items():
        assert t["category"] in CATEGORIES, name
    grouped = uv.templates_by_category()
    assert sorted(n for names in grouped.values() for n in names) == sorted(uv.TEMPLATES)
    assert list(grouped) == [k for k in CATEGORIES if k in grouped]   # catalogue order


def test_family_assignments_follow_the_classification():
    """Spectrum-resolution templates recover a spectrum (a resolution step);
    ratio-derivative ones divide then differentiate / mean-centre; ratio
    amplitude ones divide and measure amplitudes; derivative ones never divide;
    zero-order ones have no processing."""
    resolution = {"ratio_subtraction", "constant_multiplication", "extended_ratio_subtraction",
                  "factorized_recovery", "spectrum_subtraction", "constant_center"}
    for name, t in uv.TEMPLATES.items():
        ops = [s["op"] for s in t["steps"]]
        cat = t["category"]
        if cat == "zero_order":
            assert not ops, name
        elif cat == "derivative":
            assert ops == ["derivative"], name
        elif cat == "resolution":
            assert resolution & set(ops), name
        elif cat == "ratio_derivative":
            assert ops[0] in ("divide", "divide_sum") or "divide" in ops, name
            assert {"derivative", "mean_center"} & set(ops), name
        elif cat == "ratio_amplitude":
            assert ops == ["divide"], name
    for st in (s for t in uv.TEMPLATES.values() for s in t["steps"]):
        assert st["op"] in REGISTRY


def test_template_list_is_grouped_under_disabled_headings(qtbot_free_app):
    from PySide6.QtWidgets import QComboBox

    from spectro.ui.dialogs_methods import add_grouped
    cb = QComboBox()
    add_grouped(cb, uv.templates_by_category())
    model = cb.model()
    heads = [i for i in range(cb.count()) if not model.item(i).isEnabled()]
    assert len(heads) == len(uv.templates_by_category())
    assert cb.count() == len(heads) + len(uv.TEMPLATES)
    assert all(cb.itemText(i).startswith("— ") for i in heads)


@pytest.fixture(scope="module")
def qtbot_free_app():
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_method_guide_covers_every_template_and_family():
    text = GUIDE.read_text(encoding="utf-8")
    for name in uv.TEMPLATES:
        assert f"*{name}*" in text, f"template not explained in the guide: {name}"
    for title in CATEGORIES.values():
        assert f"## {title}" in text
    assert "could not be generated" not in text
    images = [line.split("](")[1].rstrip(")") for line in text.splitlines()
              if line.startswith("![")]
    assert len(images) >= 40
    for img in images:
        assert (GUIDE.parent / img).exists(), img


def test_guide_examples_recover_the_true_concentrations():
    """Every 'Result:' line of the generated guide reports recoveries within
    95–105 % (the double divisor method is approximate: within 90–110 %)."""
    import re
    text = GUIDE.read_text(encoding="utf-8")
    means = [float(m) for m in re.findall(r"mean recovery ([\d.]+) %", text)]
    assert len(means) >= 50
    assert all(90 <= m <= 110 for m in means)
    assert sum(95 <= m <= 105 for m in means) >= len(means) - 1


@pytest.mark.parametrize("tool,args", [
    ("method_guide.py", ["--only", "rd,aac"]),
    ("workflow_screenshots.py", ["--only", "special"]),
])
def test_screenshot_tools_run(tmp_path, tool, args):
    out = tmp_path / "out"
    r = subprocess.run([sys.executable, str(ROOT / "tools" / tool), "--out", str(out), *args],
                       capture_output=True, text=True, timeout=600,
                       env={**__import__("os").environ, "QT_QPA_PLATFORM": "offscreen"})
    assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
    assert list(out.glob("**/*.png")) and (out / "README.md").exists()


def test_every_method_has_a_sample_workbook_and_several_illustrations():
    import re
    text = GUIDE.read_text(encoding="utf-8")
    keys = re.findall(r'<a id="([a-z0-9_]+)"></a>', text)
    assert len(keys) == 44
    for k in keys:
        assert (GUIDE.parent / "data" / f"{k}.xlsx").exists(), k
        assert f"data/{k}.xlsx" in text
        assert (GUIDE.parent / "img" / f"{k}_concept.png").exists(), k
    # beginner sections are present for every method
    assert text.count("**In plain words.**") == 44
    assert text.count("**Checks and common mistakes.**") == 44
    images = [line for line in text.splitlines() if line.startswith("![")]
    assert len(images) >= 130


@pytest.mark.parametrize("key,compounds", [("rd", ["A", "B"]), ("aac", ["X", "Y", "Z"]),
                                           ("chemo", ["X", "Y", "Z"]), ("concval", ["A", "B"])])
def test_sample_workbooks_import_with_roles_and_concentrations(tmp_path, key, compounds):
    """A beginner's import: roles come from the sheet names, concentrations from
    the column titles; the note sheets are skipped."""
    from spectro.core.io import load_file
    from spectro.core.io.excel import role_from_sheet
    from spectro.core.naming import concentrations_from_name
    from spectro.storage.project import Project
    path = GUIDE.parent / "data" / f"{key}.xlsx"
    res = load_file(path)
    assert all(not s.metadata["sheet"].startswith("_") for s in res.spectra)
    p = Project.create(tmp_path / "b.spectro")
    for c in compounds:
        p.add_compound(c)
    roles = {i: role_from_sheet(s.metadata["sheet"]) for i, s in enumerate(res.spectra)}
    conc = {i: concentrations_from_name(s.name, compounds) for i, s in enumerate(res.spectra)}
    ids = p.import_file(path, None, roles=roles, concentrations=conc)
    recs = p.records()
    assert len(ids) == len(res.spectra)
    assert all(r.role in ("standard", "divisor", "mixture", "calibration", "sample")
               for r in recs)
    for r in recs:
        if r.role in ("standard", "mixture", "calibration"):
            assert set(compounds) <= set(r.concentrations), r.name
    p.close()


def test_workbook_calculation_sheet_reproduces_the_app():
    """The _Calculation formulas (SLOPE, INTERCEPT, found = (S − b)/a) give the
    same found values as Spectro reports in _Expected results."""
    import numpy as np
    import openpyxl
    for key in ("dw", "rd", "dd1", "rs", "zc"):
        wb = openpyxl.load_workbook(GUIDE.parent / "data" / f"{key}.xlsx")
        ws = wb["_Calculation"]
        rows = [r for r in ws.iter_rows(values_only=True)]
        start = next(i for i, r in enumerate(rows) if r[0] == "Standard") + 1
        cal = []
        for r in rows[start:]:
            if r[0] is None:
                break
            cal.append((r[1], r[2]))
        slope, intercept = np.polyfit([c for c, _ in cal], [s for _, s in cal], 1)
        mstart = next(i for i, r in enumerate(rows) if r[0] == "Mixture") + 1
        found = [(r[2] - intercept) / slope for r in rows[mstart:] if r[0]]
        exp = wb["_Expected results"]
        erows = [r for r in exp.iter_rows(values_only=True)]
        h = next(i for i, r in enumerate(erows) if r[0] == "Mixture")
        expected = [r[1] for r in erows[h + 1:]
                    if r[0] and isinstance(r[1], (int, float)) and not isinstance(r[1], bool)]
        assert np.allclose(found, expected, atol=2e-3), key

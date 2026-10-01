"""Build the illustrated method guide (docs/methods/).

For every method family and method it writes the principle, when the method
applies, how to apply it in Spectro step by step, a worked example and a
screenshot of the method being applied. The example values (wavelengths,
calibration, recoveries) are NOT typed in: the script builds a guide project
with simulated binary and ternary systems, chooses the wavelengths with the
app's finder tools, drives the real dialogs and records what they report.

    python tools/method_guide.py                 # → docs/methods/
    python tools/method_guide.py --only rd,am    # a few entries (for editing)

Runs headless (Qt "offscreen") when no display is available.
"""

from __future__ import annotations

import argparse
import html
import os
import sys
import tempfile
import traceback
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))
if not os.environ.get("DISPLAY") and sys.platform.startswith("linux"):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np  # noqa: E402
from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import (QApplication, QInputDialog, QMessageBox,  # noqa: E402
                               QTableWidgetItem)

from _png import compact_png  # noqa: E402
from guide_text import BEGINNER, PRIMER  # noqa: E402
from spectro.core import univariate as uv  # noqa: E402
from spectro.core.catalog import CATEGORIES, DESCRIPTIONS  # noqa: E402
from spectro.core.multicomponent import design_concentrations  # noqa: E402
from spectro.core.spectrum import Spectrum  # noqa: E402
from spectro.core.validation import describe  # noqa: E402

GRID = np.arange(200.0, 400.01, 0.5)
NOISE = 0.0004


def g(c, w, h=1.0):
    return h * np.exp(-0.5 * ((GRID - c) / w) ** 2)


# Unit-concentration (1 µg/mL) spectra of the simulated guide systems.
SYSTEMS = {
    "Binary A + B": {
        "about": "A absorbs below 300 nm; B overlaps A and extends alone above 300 nm "
                 "(plateau region). The spectra cross at an isoabsorptive point.",
        "pure": {"A": 0.045 * g(245, 11) + 0.012 * g(215, 7),
                 "B": 0.028 * g(268, 14) + 0.012 * g(335, 16) + 0.010 * g(222, 8)},
        "levels": (4, 8, 12, 16, 20),
        "mixtures": [{"A": 6, "B": 10}, {"A": 10, "B": 6}, {"A": 14, "B": 14},
                     {"A": 8, "B": 16}, {"A": 18, "B": 5}],
        "divisors": {"B": 10},
    },
    "Ternary X + Y + Z": {
        "about": "X and Y overlap completely below 300 nm; Z overlaps both and extends alone "
                 "above 330 nm.",
        "pure": {"X": 0.040 * g(240, 13) + 0.022 * g(275, 9),
                 "Y": 0.030 * g(255, 12) + 0.018 * g(290, 8),
                 "Z": 0.024 * g(258, 14) + 0.010 * g(355, 22)},
        "levels": (4, 8, 12, 16, 20, 24),
        "mixtures": [{"X": 10, "Y": 10, "Z": 10}, {"X": 20, "Y": 10, "Z": 10},
                     {"X": 6, "Y": 6, "Z": 18}, {"X": 10, "Y": 20, "Z": 10},
                     {"X": 12, "Y": 8, "Z": 6}],
        "divisors": {"X": 12, "Y": 12, "Z": 24},
    },
    "Ternary U + V + W": {
        "about": "Successive extension: U absorbs alone above 320 nm, U and V together at "
                 "275–300 nm, all three below 260 nm.",
        "pure": {"U": 0.030 * g(300, 14) + 0.020 * g(325, 10),
                 "V": 0.025 * g(268, 9) + 0.012 * g(232, 8),
                 "W": 0.045 * g(246, 7)},
        "levels": (2, 6, 10, 14, 18),
        "mixtures": [{"U": 10, "V": 10, "W": 10}, {"U": 16, "V": 4, "W": 4},
                     {"U": 4, "V": 14, "W": 8}, {"U": 8, "V": 8, "W": 16},
                     {"U": 12, "V": 6, "W": 12}],
        "divisors": {"U": 10, "V": 8},
    },
}
SYSTEM_OF = {c: name for name, s in SYSTEMS.items() for c in s["pure"]}


# --------------------------------------------------------------------------- #
# Guide content
# --------------------------------------------------------------------------- #
@dataclass
class Entry:
    key: str
    title: str
    category: str
    principle: str
    when: str
    steps: list[str]
    example: str = ""                        # name of the Guide method that runs it
    refs: str = ""
    notes: str = ""
    result: dict = field(default_factory=dict)


E = Entry
ENTRIES: list[Entry] = [
    # ------------------------------------------------------------- zero order
    E("direct", "Direct measurement at λmax", "zero_order",
      "A = a·b·C (Beer–Lambert): the absorbance at a wavelength where only the analyte "
      "absorbs is proportional to its concentration.",
      "The analyte has a region where no other compound absorbs (usually the extended "
      "component of a mixture).",
      ["Methods → Univariate calibration → template *Direct (zero order, λmax)*.",
       "Choose the compound, then drag the λ marker to a wavelength where only it absorbs "
       "(check the pure spectra of the other compounds).",
       "Check the calibration standards and the mixtures → *Calibrate* → *Determine*."],
      "direct"),
    E("dw", "Dual wavelength (DW)", "zero_order",
      "ΔA = A(λ1) − A(λ2). At λ1 and λ2 the interferent Y has the same absorbance, so its "
      "contribution cancels and ΔA ∝ C_X.",
      "Y has two wavelengths with equal absorbance where X absorbs differently.",
      ["Tools → Spectral finder → *Equal amplitude*: select a pure Y spectrum and λ1 (usually "
       "λmax of X) to list the λ2 values where Y is equal.",
       "Methods → Univariate calibration → template *Dual wavelength*; enter λ1 and λ2.",
       "Calibrate on pure X standards (ΔA vs C), determine the mixtures."],
      "dw", "EI18"),
    E("idw", "Induced dual wavelength (IDW)", "zero_order",
      "ΔA = A(λ1) − F·A(λ2) with the equality factor F = A_Y(λ1)/A_Y(λ2) of pure Y. Any "
      "λ pair works: F 'induces' equal Y absorbance.",
      "Y has no pair of equal absorbance in a useful region, or a better sensitivity for X is "
      "obtained at another pair.",
      ["Template *Induced dual wavelength*; enter λ1 and λ2.",
       "In the measurement, choose a pure Y standard as *Interferent spectrum for F* (F is "
       "computed from it) or type F.",
       "Calibrate on pure X standards, determine."],
      "idw", "EI18"),
    E("auc1", "Area under the curve — single component", "zero_order",
      "The integral of the absorbance between λ1 and λ2 is proportional to C.",
      "The analyte absorbs alone over a band; integration averages out noise.",
      ["Template *Area under curve (single component)*; set the integration range.",
       "Choose a zero or linear baseline for the area. Calibrate, determine."],
      "auc1"),
    E("q", "Absorbance ratio (Q-analysis)", "zero_order",
      "At the isoabsorptive point A_iso = a_iso(C_X + C_Y); the ratio Q = A(λ2)/A(iso) of "
      "the mixture interpolates between those of X and Y: "
      "C_X = (Q_m − Q_Y)/(Q_X − Q_Y) · A_iso/a_iso.",
      "Binary mixture with an isoabsorptive point.",
      ["Methods → Binary two-signal methods → tab *Zero order: Q-analysis*.",
       "Choose X and Y, press *Find isoabsorptive point*, set λ2 (λmax of X).",
       "Check the X standards, Y standards and the mixtures → *Calculate*."],
      "q"),
    E("as", "Absorbance subtraction (AS)", "zero_order",
      "X absorbs alone at λ2. The amplitude factor AF = A_X(iso)/A_X(λ2) of pure X turns "
      "the mixture absorbance at λ2 into X's absorbance at the isoabsorptive point; the "
      "total at iso gives C_X + C_Y.",
      "Binary mixture with an isoabsorptive point and a region where one compound "
      "absorbs alone (the extended one).",
      ["Binary two-signal methods → *Zero order: absorbance subtraction*.",
       "X = the compound that absorbs alone at λ2; iso = isoabsorptive point.",
       "Check pure X standards (for AF), standards for the iso calibration, mixtures → "
       "*Calculate*."],
      "as_"),
    E("aas", "Advanced absorbance subtraction (AAS)", "zero_order",
      "λ1 and λ2 are chosen where Y has equal absorbance (one of them at the isoabsorptive "
      "point): ΔA = A(λ2) − A(λ1) gives C_X; X's absorbance at λ2 follows from ΔA and is "
      "subtracted; the remainder gives C_Y.",
      "Binary mixture with an isoabsorptive point and a second wavelength where Y equals its "
      "iso absorbance.",
      ["Binary two-signal methods → *Zero order: advanced absorbance subtraction*.",
       "λ2 = isoabsorptive point; press *Find λ1 where Y equals its value at λ2*.",
       "Check X standards, Y standards and mixtures → *Calculate*."],
      "aas", "EI18"),
    E("mafm", "Successive absorption factor method (MAFM)", "zero_order",
      "Compounds are taken in order of extension. The first absorbs alone at λ1. Its "
      "absorption factors F(λ) = A(λ)/A(λ1) give its contribution at the next wavelengths, "
      "which is subtracted; the second compound is then alone at λ2, and so on.",
      "Successive extension: one compound alone at long λ, the next one overlapped only by "
      "the first, … (binary: the classic absorption factor method).",
      ["Methods → Progressive resolution → tab *Zero order: successive absorption factor*.",
       "Enter the compounds in order with the wavelength where each is 'added' (the first "
       "alone, the second overlapped only by the first …); optionally a quantitation λ.",
       "Check the pure standards and the mixtures → *Calculate*; *Save method* to reuse it."],
      "mafm", "AW17"),
    # ------------------------------------------------------------- equations
    E("vierordt", "Simultaneous equations (Vierordt)", "equations",
      "A(λ1) = a_X1·C_X + a_Y1·C_Y and A(λ2) = a_X2·C_X + a_Y2·C_Y, solved for both "
      "concentrations (more wavelengths → least squares).",
      "Each compound has a band where it dominates; check the condition number of K.",
      ["Methods → Equation methods. Tick the compounds, enter one amplitude signal per "
       "compound (their λmax) or more.",
       "Calibrate on the pure standards (*Fit*) — the absorptivity matrix K and its condition "
       "number are shown — then *Determine*."],
      "vierordt"),
    E("bivariate", "Bivariate calibration (Kaiser)", "equations",
      "Two wavelengths that maximise the determinant of the sensitivity matrix (Kaiser) are "
      "used with linear calibrations that include intercepts.",
      "Binary mixtures; picks the best-conditioned wavelength pair automatically.",
      ["Equation methods → press *Kaiser* (uses one pure standard of each compound) — the "
       "best pair is filled in and intercepts are switched on.",
       "*Fit* on the pure standards, *Determine*."],
      "bivariate"),
    E("auc_eq", "Area under the curve — simultaneous (binary / ternary)", "equations",
      "The areas in n ranges are linear in the n concentrations: A_k = Σ a_kj·C_j; solved "
      "with Cramer's rule / least squares.",
      "Overlapping bands; areas are less noisy than single amplitudes.",
      ["Equation methods → one *area* signal per compound (λ1–λ2 ranges chosen where the "
       "compounds differ most). *Fit*, *Determine*."],
      "auc_eq", "AW17, SA14"),
    # ------------------------------------------------------------- derivative
    E("zc", "Zero-crossing derivative (D1–D4)", "derivative",
      "The derivative dⁿA/dλⁿ of Y is zero at its zero-crossing wavelengths; there the "
      "mixture's derivative depends on X only.",
      "Y's derivative crosses zero where X's derivative is large.",
      ["Tools → Spectral finder → *Zero-crossings* on a processed (D1) pure Y spectrum, or "
       "watch the processed plot in the Univariate dialog.",
       "Template *Zero-crossing derivative (D1)* (or *Zero-crossing derivative (D2)*); set "
       "Δλ (4 nm) and the scaling factor (10) in the Derivative step; put λ at Y's "
       "zero-crossing.",
       "Calibrate on pure X, determine."],
      "zc"),
    E("p2p", "Derivative peak-to-peak", "derivative",
      "The difference between the maximum and minimum of the derivative in a range is "
      "proportional to C (larger signal, insensitive to baseline offsets).",
      "The analyte's derivative band is free from the other compounds in that range.",
      ["Template *Derivative peak-to-peak*; set the range around the derivative band.",
       "Calibrate, determine."],
      "p2p"),
    E("d1dwl", "Dual wavelength in derivative mode (D1 DWL)", "derivative",
      "ΔD1 = D1(λ1) − D1(λ2) at two wavelengths where Y's D1 is equal → depends on X only.",
      "Y's derivative has two equal amplitudes; avoids divisors (no ratio spectra).",
      ["Finder → *Equal amplitude* on the D1 spectrum of pure Y.",
       "Template *Dual wavelength in derivative mode*; enter λ1, λ2. Calibrate, determine."],
      "d1dwl", "LO15"),
    # ------------------------------------------------------------- ratio derivative
    E("dd1", "Derivative ratio (DD1 / DR1)", "ratio_derivative",
      "(X + Y)/Y′ = X/Y′ + constant; the first derivative removes the constant and leaves "
      "d(X/Y′)/dλ ∝ C_X, read at a maximum or minimum.",
      "Binary mixtures; the classic ratio-spectra method.",
      ["Template *Derivative ratio (DD1)*; choose the divisor Y′ in the Divide step "
       "(a standard of Y, or a unit-concentration spectrum — see *Normalize → concentration*).",
       "Set Δλ and scaling in the Derivative step, put λ at a peak of the processed X spectra.",
       "Calibrate on pure X (processed the same way), determine."],
      "dd1", "EI18"),
    E("d1dr", "Derivative ratio of derivative spectra (D1 DR)", "ratio_derivative",
      "D1(mixture) ÷ D1(Y′) = D1X/D1Y′ + constant; differentiating again removes Y.",
      "When the ordinary ratio spectra give no usable peaks for X.",
      ["Template *Derivative ratio of D1 spectra (D1 DR)*: Derivative → Divide (with "
       "*Differentiate divisor first* = 1) → Derivative.",
       "Choose the divisor Y′, set λ at a peak away from the divisor's D1 zero-crossings.",
       "Calibrate, determine."],
      "d1dr", "LO15"),
    E("mcr", "Mean centering of ratio spectra (MCR)", "ratio_derivative",
      "The ratio spectrum X/Y′ + constant is mean centred over a range: the constant "
      "vanishes, MC(X/Y′) ∝ C_X.",
      "Binary mixtures; no derivative, so less noise amplification.",
      ["Template *Mean centering of ratio spectra (MCR)*; choose Y′; set the mean-centring "
       "range inside the region where Y′ is not near zero.",
       "Read at a maximum/minimum of the processed X spectra. Calibrate, determine."],
      "mcr", "EI18"),
    E("mcr3", "Mean centering of ratio spectra — ternary", "ratio_derivative",
      "÷ Z′ → MC → ÷ MC(Y′/Z′) → MC: after the first division Z is a constant (removed by "
      "MC); after dividing by the mean-centred ratio of Y, Y is a constant too (removed by "
      "the second MC). The result depends on X only.",
      "Ternary mixtures; two divisors, one wavelength.",
      ["Template *Mean centering of ratio spectra (ternary)*: in the first Divide choose Z′, "
       "in *Divide by mean-centred ratio* choose Y′ (second component) and Z′ (same first "
       "divisor), same range in all mean-centring steps.",
       "Read X at an extremum away from the spikes where MC(Y′/Z′) crosses zero. Calibrate, "
       "determine."],
      "mcr3", "SA14"),
    E("sdr", "Successive derivative ratio (ternary)", "ratio_derivative",
      "(X+Y+Z)/Y′ → D1 removes Y; dividing by D1(Z′/Y′) makes Z a constant; a second D1 "
      "removes it.",
      "Ternary mixtures where X has peaks in the final spectrum.",
      ["Create the second divisor once: Process → pipeline on a Z standard: Divide by Y′ → "
       "Derivative (same Δλ) → saved as a derived spectrum.",
       "Template *Successive derivative ratio (ternary)*: first Divide = Y′, second Divide = "
       "that derived D1(Z′/Y′) spectrum.",
       "Read X at a peak away from spikes. Calibrate, determine."],
      "sdr"),
    E("dd", "Double divisor ratio derivative (ternary)", "ratio_derivative",
      "The mixture is divided by the sum of the other two standards (Y′ + Z′) and "
      "differentiated. Y and Z cancel only when they are present in the same proportion as "
      "in the divisor, so the method is approximate.",
      "Ternary mixtures in a fixed ratio (e.g. one dosage form); check recoveries on "
      "laboratory mixtures of varying ratio.",
      ["Template *Double divisor ratio derivative (ternary)*; choose Y′ and Z′ (equal "
       "concentrations, as in the papers) in *Divide by sum*.",
       "Read X at a D1 extremum, calibrate, determine."],
      "dd", "SA14"),
    # ------------------------------------------------------------- ratio amplitude
    E("rd", "Ratio difference (RD)", "ratio_amplitude",
      "ΔP = P(λ1) − P(λ2) of the ratio spectrum X/Y′ + constant: the constant cancels, "
      "ΔP ∝ C_X.",
      "Binary mixtures; two wavelengths with a large difference for X/Y′.",
      ["Template *Ratio difference (RD)*; choose Y′ (a standard or a unit-concentration "
       "spectrum).",
       "Drag λ1 and λ2 to a peak and a trough of the processed X spectra. Calibrate, "
       "determine."],
      "rd", "EI18, LO15"),
    E("dad", "Dual amplitude difference (ternary)", "ratio_amplitude",
      "Divide by Z′ (Z becomes a constant). Choose λ1, λ2 where Y/Z′ has equal amplitudes: "
      "P(λ1) − P(λ2) cancels Y and Z together and depends on X only (an equality factor F "
      "from Y corrects a residual difference).",
      "Ternary mixtures.",
      ["Finder → *Equal amplitude* on the ratio spectrum Y′/Z′ (process a Y standard by "
       "Divide by Z′ first).",
       "Template *Dual amplitude difference (ternary)*: divisor Z′, λ1 and λ2, interferent Y′ "
       "for F.",
       "Calibrate on pure X, determine."],
      "dad"),
    E("cv", "Constant value (ratio plateau)", "ratio_amplitude",
      "Where only Y absorbs, the ratio spectrum (X + Y)/Y′ is a flat line (plateau) equal to "
      "C_Y/C_Y′.",
      "Y extends alone (the extended component).",
      ["Tools → Spectral finder → *Plateaus* on a mixture ÷ Y′ to find the flat region.",
       "Template *Ratio plateau / constant value (Y)*; divisor Y′; set the plateau range.",
       "Calibrate on pure Y standards, determine."],
      "cv", "LO15, FA21"),
    E("concval", "Concentration value (no regression)", "ratio_amplitude",
      "With a unit-concentration (normalized) divisor the plateau value IS the concentration "
      "of Y — no calibration line is needed.",
      "Same conditions as constant value.",
      ["Make the normalized divisor once: Process → *Normalize* (mode *concentration*) on a "
       "pure Y standard.",
       "Template *Concentration value*; choose that unit spectrum; the option "
       "*Concentration value: signal = concentration* is ticked.",
       "*Calibrate* (still reports the calibration line as a check), *Determine*."],
      "concval", "FA21"),
    E("am", "Amplitude modulation (AM)", "ratio_amplitude",
      "With a unit-concentration divisor Y′ the ratio spectrum at the isoabsorptive point "
      "equals C_X + C_Y and the plateau equals C_Y; C_X = total − C_Y. A unified regression "
      "at the iso point corrects small deviations.",
      "Binary mixture with an isoabsorptive point; Y extended.",
      ["Make the unit Y spectrum (Normalize → concentration).",
       "Binary two-signal methods → *Ratio: amplitude modulation*: X, Y, iso, plateau, "
       "divisor = unit Y, divisor concentration = 1.",
       "Check standards of both compounds (unified calibration) and mixtures → *Calculate*."],
      "am", "ZA20, LO15"),
    E("iam", "Induced amplitude modulation (IAM)", "ratio_amplitude",
      "Amplitude modulation without an isoabsorptive point: P(λ) − C_Y = r·C_X with "
      "r = (a_X/a_Y)(λ) from pure X.",
      "Y extended; no isoabsorptive point needed.",
      ["Binary two-signal methods → *Ratio: induced amplitude modulation*; λ for X, plateau.",
       "Check X standards, Y standards (the unit divisor is averaged from them), mixtures."],
      "iam"),
    E("cvad", "Constant value via amplitude difference (CV-AD)", "ratio_amplitude",
      "Divide by Y′. The ratio difference of X between λ1 and λ2 gives X's postulated "
      "amplitude at λ2 (from a line ΔP vs P(λ2) of pure X); the recorded amplitude minus it "
      "is the constant C_Y/C_Y′.",
      "Binary (or ternary where the third compound does not absorb there); no plateau needed.",
      ["Methods → Progressive resolution → *Ratio: amplitude centering*.",
       "Compounds X and Y; divisor Y′; divisor compound Y; common λc = λ2; amplitude "
       "difference row: X, λ1, λ2; amplitude subtraction for Y.",
       "Check standards and mixtures → *Calculate*."],
      "cvad", "LO15"),
    E("aac", "Advanced amplitude centering — partial overlap (AAC)", "ratio_amplitude",
      "One divisor Z′ and one wavelength λc for all compounds: the plateau gives Z; after "
      "subtracting it, P(λ1) − F_Y·P(λ2) depends on X only (equality factor of Y) and gives "
      "X's postulated amplitude at λc; Y = recorded − Z − X at λc.",
      "Ternary; Z extended alone (plateau).",
      ["Progressive resolution → *Ratio: amplitude centering*: divisor Z′, divisor compound Z, "
       "*Plateau* ticked with its range, λc.",
       "Amplitude difference row: X, λc, λ2, factor from Y. Amplitude subtraction for Y.",
       "*Calculate*; *Save method* to reuse it. The optimizer proposes these settings "
       "automatically."],
      "aac", "SL17"),
    E("macm", "Amplitude centering — complete overlap (AAC / MACM)", "ratio_amplitude",
      "No plateau: X from a λ pair where Y/Z′ is equal, Y from a λ pair where X/Z′ is equal "
      "(Z/Z′ is constant and cancels in both); Z by subtraction at the common λc.",
      "Ternary with severe overlap; one divisor.",
      ["Any of the three compounds can be the divisor (called Z here); the method "
       "optimizer tries each one and proposes λc and the λ pairs.",
       "By hand: Finder → *Equal amplitude* on Y/Z′ and on X/Z′ at the chosen λc.",
       "Amplitude centering: divisor Z′, divisor compound Z, no plateau; rows X (λc, partner "
       "of Y) and Y (λc, partner of X); subtraction for Z.",
       "*Calculate*."],
      "macm", "SL17, AW17"),
    E("ridss", "Ratio difference–isoabsorptive (RIDSS)", "ratio_amplitude",
      "Z from the plateau; Y from a ratio difference at a pair where X/Z′ is equal; at the "
      "isoabsorptive point of X and Y the amplitude after removing Z is C_X + C_Y (unified "
      "regression), so X = total − Y.",
      "Ternary; Z extended; X and Y have an isoabsorptive point.",
      ["Finder → *Isoabsorptive points* (X and Y standards) for λc, *Equal amplitude* on X/Z′ "
       "for the partner.",
       "Amplitude centering: plateau, row Y (λc, partner), subtraction for X, *Unified "
       "regression* ticked."],
      "ridss", "AM14"),
    # ------------------------------------------------------------- resolution
    E("rs", "Ratio subtraction (RS) / spectrum subtraction of the extended component",
      "resolution",
      "(X + Y)/Y′ − constant (plateau) = X/Y′; × Y′ gives X's zero-order spectrum, read at "
      "X's λmax. Spectro computes it as mixture − constant·Y′ (identical, exact everywhere).",
      "Y extended alone (plateau).",
      ["Template *Ratio subtraction / spectrum subtraction (X)*; divisor Y′, plateau range.",
       "The processed mixtures now look like pure X: read at X's λmax. Calibrate, determine."],
      "rs", "ZA20, LO15"),
    E("cm", "Constant multiplication (CM) / SS-CM", "resolution",
      "The plateau constant × Y′ is Y's zero-order spectrum in the mixture (CM); subtracting "
      "it from the mixture gives X (spectrum subtraction, SS). Both spectra are then read at "
      "their λmax.",
      "Y extended alone.",
      ["Template *Constant multiplication / SS-CM (extended Y)*; divisor Y′, plateau.",
       "Read Y at its λmax. For X use the ratio subtraction template (same divisor/plateau)."],
      "cm", "ZA20, FA21"),
    E("ers", "Extended ratio subtraction (ERS)", "resolution",
      "After ratio subtraction recovers X, X's amount is matched to its pure spectrum X′ and "
      "subtracted from the mixture → Y's spectrum, read at Y's λmax (also where Y has no "
      "region of its own).",
      "Y extended; Y measured at its band maximum.",
      ["Template *Extended ratio subtraction (Y)*; divisor Y′, pure X′, plateau. Read Y at λmax."],
      "ers", "LO15, EM18"),
    E("srs", "Successive ratio subtraction (SRS, ternary)", "resolution",
      "Ratio subtraction with the most extended component's divisor removes it; a second "
      "ratio subtraction with the next component's divisor (its own plateau) removes that "
      "one; the third component's spectrum remains.",
      "Successive extension (U alone at long λ, V alone after U is removed).",
      ["Template *Successive ratio subtraction (ternary SRS, Z)*: first step divisor U′ with "
       "U's plateau, second step divisor V′ with V's plateau.",
       "Read W at its λmax. Calibrate, determine."],
      "srs", "EM18"),
    E("sss", "Successive spectrum subtraction (ternary)", "resolution",
      "Spectrum subtraction (factor k from a wavelength where the first component absorbs "
      "alone) removes it; ratio subtraction then removes the second.",
      "Successive extension.",
      ["Template *Successive spectrum subtraction (ternary)*: Spectrum subtraction with U′ "
       "at a λ where only U absorbs, then Ratio subtraction with V′ and V's plateau.",
       "Read W at its λmax."],
      "sss"),
    E("ss", "Spectrum subtraction (SS)", "resolution",
      "k = A_mixture(λ)/A_Y′(λ) at a wavelength (or plateau range) where only Y absorbs; "
      "mixture − k·Y′ = X.",
      "Y absorbs alone at some wavelength.",
      ["Template *Spectrum subtraction*: reference Y′, wavelength (and optionally *…to* for a "
       "range), derivative order 0. Read X at λmax."],
      "ss"),
    E("fzm", "Factorized zero-order method (FZM)", "resolution",
      "k from the D1 amplitude at a zero-crossing of the other component: "
      "k = D1_mixture(λ)/D1_X′(λ); k·X′ is X's zero-order spectrum, read at its λmax.",
      "The other component's derivative crosses zero where X's is large.",
      ["Template *Factorized zero-order (FZM)*: reference X′, wavelength = zero-crossing of "
       "Y's D1, derivative order 1. Read X at λmax."],
      "fzm"),
    E("cc", "Constant center (CC)", "resolution",
      "For X the ratio r = P_X(λ1)/P_X(λ2) of X/Y′ is constant, so the constant of Y is "
      "k = P2 − (P1 − P2)/(r − 1); (P − k)·Y′ = X, k·Y′ = Y.",
      "No plateau needed; X's ratio spectrum has different amplitudes at λ1 and λ2.",
      ["Template *Constant center*: divisor Y′, pure X′, λ1, λ2, recover X (or Y).",
       "Read the recovered component at its λmax."],
      "cc", "LO15"),
    E("dt", "Derivative transformation (DT, DT-SS)", "resolution",
      "In the D1 spectra the interferent is flat (D1 = 0) over a region where Y's D1 is not: "
      "the mean of D1(mixture)/D1(Y′) there is k; k·Y′ (zero order) is Y's spectrum; "
      "mixture − k·Y′ (DT-SS) is the other's.",
      "The interferent is flat (or zero) over a region where Y still changes.",
      ["Template *Derivative transformation (DT, recover zero order)*: reference Y′, "
       "wavelength and *…to* = the plateau of the D1 ratio, derivative order 1.",
       "Read Y at its λmax (zero order)."],
      "dt", "FA21"),
    E("ds", "Derivative subtraction (DS) and DS-CM", "resolution",
      "Ratio subtraction / constant multiplication in the D1 domain: D1 of the mixture ÷ D1 "
      "of Y′ has a plateau where only Y's D1 is non-zero; DS gives X's D1 spectrum, DS-CM "
      "gives Y's.",
      "Useful when the zero-order plateau is noisy (small extension).",
      ["Template *Derivative subtraction (DS, X in D1)* or *…(DS-CM, Y in D1)*: Derivative, "
       "then Ratio subtraction / Constant multiplication with *Differentiate divisor first* "
       "= 1 and the D1 plateau.",
       "Read the component's D1 at a peak. Calibrate, determine."],
      "ds", "ZA20, LO15"),
    # ------------------------------------------------------------- multivariate
    E("chemo", "Chemometrics: CLS, ILS, PCR, PLS, MCR-ALS, ANN, SVR", "multivariate",
      "Calibration on whole spectra. CLS uses pure spectra (K matrix); ILS/PCR/PLS regress "
      "concentrations on spectra of a mixture training set; MCR-ALS resolves pure profiles; "
      "ANN and SVR model non-linearity.",
      "Severe overlap, three or more compounds, or when univariate methods fail. Needs a "
      "designed training set (Tools → Calibration design, e.g. Brereton 5-level).",
      ["Tools → Calibration design to plan the training mixtures; import them with role "
       "*calibration*.",
       "Methods → Chemometrics: model type, wavelength range, preprocessing.",
       "*Cross-validate* (number of components is suggested), *Fit + predict* the "
       "validation mixtures; check T²/Q outliers; *Save method*."],
      "chemo", "SL17"),
    # ------------------------------------------------------------- standard addition
    E("stdadd", "Standard addition (recovery in the matrix)", "standard_addition",
      "Known amounts of pure analyte are added to the sample; the found − unspiked amount "
      "divided by the added amount is the recovery; the x-intercept of found vs added is the "
      "sample content.",
      "Check accuracy in a real matrix (excipients) with any calibrated method.",
      ["Univariate calibration with a calibrated method; import the unspiked and spiked "
       "samples with the analyte concentration = amount ADDED (0 for the unspiked).",
       "Check them in *Spectra to determine* → *Standard addition…*."],
      "stdadd"),
    E("hpsam", "H-point standard addition (HPSAM)", "standard_addition",
      "Standard additions of X measured at λ1 and λ2 where Y has equal absorbance; the two "
      "addition lines meet at H(−C_X, A_Y): the analyte concentration free of the "
      "interferent and the matrix.",
      "Binary; Y has two wavelengths with equal absorbance.",
      ["Binary two-signal methods → *Standard addition: H-point*: analyte, λ1, λ2.",
       "Check the sample and its additions (X concentration = amount added) → *Calculate*."],
      "hpsam"),
    E("enrich", "Sample enrichment: spiking and spectrum addition", "standard_addition",
      "A minor component below its linear range is raised by adding a known amount of pure "
      "standard (spiking) or by adding the stored spectrum of a standard (spectrum addition); "
      "the added amount is subtracted from the result.",
      "Dosage forms with a very unequal ratio (e.g. 20 : 1).",
      ["Spiking: Univariate calibration → *Enrichment added* = amount added → the results get "
       "a *Found − added* column.",
       "Spectrum addition: Process → *Add spectrum* (a stored standard) on the sample, then "
       "determine and subtract the added amount."],
      "enrich", "ZA20, LO15, FA21"),
]

REFS = {
    "AM14": "Abdelrahman et al., Anal. Methods 6 (2014) 509, doi:10.1039/c3ay41564c",
    "SA14": "Abdelrahman et al., Spectrochim. Acta A 124 (2014) 389, "
            "doi:10.1016/j.saa.2014.01.020",
    "LO15": "Lotfy et al., Spectrochim. Acta A 136 (2015) 937, doi:10.1016/j.saa.2014.09.117",
    "AW17": "Abdelwahab & Mohamed, Chem. Pharm. Bull. 65 (2017) 558",
    "SL17": "Saleh et al., Int. J. Pharm. Pharm. Sci. 9 (2017) 43, "
            "doi:10.22159/ijpps.2017v9i5.16960",
    "EI18": "Eissa & Abou Al Alamein, Spectrochim. Acta A 193 (2018) 365, "
            "doi:10.1016/j.saa.2017.12.050",
    "EM18": "Emam et al., Spectrochim. Acta A (2018), doi:10.1016/j.saa.2017.11.034",
    "ZA20": "Zaghary et al., J. Anal. Chem. 75 (2020) 742, doi:10.1134/S1061934820060180",
    "FA21": "Fahmy et al., Spectrochim. Acta A 261 (2021) 119999, "
            "doi:10.1016/j.saa.2021.119999",
}


# --------------------------------------------------------------------------- #
# Running the examples in the real dialogs
# --------------------------------------------------------------------------- #
class Failed(Exception):
    pass


def _raise(parent, exc, title="Spectro"):
    raise Failed(str(exc))


def check_only(lst, ids):
    for i in range(lst.count()):
        it = lst.item(i)
        it.setCheckState(Qt.Checked if it.data(Qt.UserRole) in ids else Qt.Unchecked)


def fmt(v):
    return f"{v:.4g}" if isinstance(v, float) else str(v)


class Guide:
    def __init__(self, out: Path, workdir: Path, app: QApplication):
        import spectro.ui.dialogs_data as dd
        import spectro.ui.dialogs_methods as dm
        import spectro.ui.dialogs_tools as dt
        import spectro.ui.main_window as mw
        import spectro.ui.widgets as w
        for mod in (w, dd, dm, dt, mw):
            mod.error = _raise
            setattr(mod, "ask_reason", lambda *a, **k: "guide")
        QMessageBox.information = staticmethod(lambda *a, **k: None)
        QInputDialog.getText = staticmethod(lambda *a, **k: ("guide example", True))
        self.out, self.app = out, app
        out.mkdir(parents=True, exist_ok=True)
        (out / "img").mkdir(exist_ok=True)
        (out / "data").mkdir(exist_ok=True)
        self.win = mw.MainWindow()
        self.win.resize(1400, 860)
        from spectro.storage.project import Project
        self.p = Project.create(workdir / "guide.spectro", "Method guide")
        self.build()
        self.win._attach(self.p)

    # ------------------------------------------------------------ project
    def build(self):
        p = self.p
        rng_seed = 0
        self.ids: dict[str, int] = {}
        self.mix: dict[str, list[int]] = {}
        self.trial: dict[str, int] = {}
        for name, sysd in SYSTEMS.items():
            for c in sysd["pure"]:
                p.add_compound(c)
            tid = self.trial[name] = p.add_trial(name, sysd["about"])
            comps = list(sysd["pure"])

            def spec(conc, label, noise=NOISE):
                nonlocal rng_seed
                rng_seed += 1
                y = sum(sysd["pure"][k] * v for k, v in conc.items())
                y = y + np.random.default_rng(rng_seed).normal(0, noise, GRID.size)
                return Spectrum(GRID, y, name=label,
                                concentrations={k: float(conc.get(k, 0.0)) for k in comps})
            for c in comps:
                for v in sysd["levels"]:
                    self.ids[f"{c} {v}"] = p.add_spectrum(spec({c: v}, f"{c} {v} µg/mL"),
                                                          tid, role="standard")
            for c, v in sysd["divisors"].items():
                self.ids[f"{c}′"] = p.add_spectrum(spec({c: v}, f"{c}′ {v} µg/mL (divisor)",
                                                        noise=NOISE / 3), tid, role="divisor")
                self.ids[f"{c} unit"] = p.process([self.ids[f"{c}′"]],
                                                  [{"op": "normalize",
                                                    "params": {"mode": "concentration"}}],
                                                  "unit concentration")[0]
            self.mix[name] = [p.add_spectrum(spec(m, "Mix " + ", ".join(
                f"{k} {v:g}" for k, v in m.items())), tid, role="mixture")
                for m in sysd["mixtures"]]
        # spectra processed once for the finder screenshots
        self.ids["B′ D1"] = p.process([self.ids["B′"]], [{"op": "derivative",
                                                          "params": dict(self.D1)}], "D1")[0]
        self.ids["Mix ÷ B′"] = p.process([self.mix["Binary A + B"][0]],
                                         [{"op": "divide",
                                           "params": {"reference": self.ids["B′"]}}],
                                         "ratio to B′")[0]
        # second divisor for the successive derivative ratio: D1(Z′ ÷ Y′)
        self.ids["D1(Z′/Y′)"] = p.process(
            [self.ids["Z′"]], [{"op": "divide", "params": {"reference": self.ids["Y′"]}},
                               {"op": "derivative", "params": {"order": 1,
                                                               "delta_lambda": 4.0}}],
            "D1 of ratio to Y′")[0]
        # chemometric training set for X, Y, Z (Brereton design)
        tern = SYSTEMS["Ternary X + Y + Z"]
        tid = self.trial["Ternary X + Y + Z"]
        for i, c in enumerate(design_concentrations({"X": 14, "Y": 14, "Z": 14},
                                                    {"X": 4, "Y": 4, "Z": 4})):
            y = sum(tern["pure"][k] * v for k, v in c.items())
            y = y + np.random.default_rng(900 + i).normal(0, NOISE, GRID.size)
            p.add_spectrum(Spectrum(GRID, y, name=f"Training {i + 1}", concentrations=c), tid,
                           role="calibration")

    def unit(self, c) -> Spectrum:
        return Spectrum(GRID, SYSTEMS[SYSTEM_OF[c]]["pure"][c], name=c)

    # ------------------------------------------------------------ helpers
    def shot(self, widget, key, height=820, keep=False) -> str:
        sz = widget.size()          # (FinderDialog has a 'width' attribute)
        widget.resize(max(sz.width(), 1300), max(sz.height(), height))
        widget.show()
        for _ in range(4):
            self.app.processEvents()
        pm = widget.grab()
        if pm.width() > 1200:
            pm = pm.scaledToWidth(1200, Qt.SmoothTransformation)
        name = f"img/{key}.png"
        pm.save(str(self.out / name))
        compact_png(self.out / name)
        if not keep:
            widget.close()
        return name

    def table(self, comps, spectra_names, found, taken):
        rows, recs = [], {c: [] for c in comps}
        for n, f, t in zip(spectra_names, found, taken):
            row = [n]
            for c, fv, tv in zip(comps, f, t):
                r = 100 * fv / tv if tv else None
                if r is not None:
                    recs[c].append(r)
                row += [f"{fv:.3f}", "" if tv is None else f"{tv:g}",
                        "" if r is None else f"{r:.2f}"]
            rows.append(row)
        head = ["Mixture"] + [h for c in comps for h in (f"{c} found", f"{c} taken", "Rec. %")]
        summ = []
        for c, v in recs.items():
            if len(v) > 1:
                d = describe(v)
                summ.append(f"{c}: mean recovery {d['mean']:.2f} %, RSD {d['rsd']:.2f} %")
        return {"headers": head, "rows": rows, "summary": summ}

    def univariate(self, key, compound, template, refs=None, params=None, meas=None,
                   direct=False, mixtures=None, notes=None, spike=0.0):
        """Apply a template in the Univariate dialog exactly as a user would."""
        from spectro.ui.dialogs_methods import UnivariateDialog
        d = UnivariateDialog(self.win)
        d.resize(1400, 860)
        d.compound.setCurrentText(compound)
        d.template.setCurrentText(template)
        steps = d.pipe.get_steps()
        for i, kv in (refs or {}).items():
            for k, v in kv.items():
                steps[i]["params"][k] = self.ids[v] if isinstance(v, str) else v
        for i, kv in (params or {}).items():
            steps[i]["params"].update(kv)
        d.pipe.set_steps(steps)
        if meas:
            m = d.meas.get()
            m["params"].update(meas)
            d.meas.set(m)
        d.direct.setChecked(direct)
        d.spike.setValue(spike)
        mix = mixtures if mixtures is not None else self.mix[SYSTEM_OF[compound]]
        check_only(d.test, mix)
        d._calibrate()
        d._predict()
        reg = d.method.regression
        d.tabs.setCurrentIndex(0)
        from spectro.core.operations import apply_pipeline
        res_ = self.p.resolver()
        d.plot.plot_spectra([apply_pipeline(s, d.method.steps, res_)
                             for s in self.p.spectra(mix)], keep_range=False)
        d._markers()
        images = [self.concept_univariate(key, compound, d.method)]
        images.append((self.shot(d, f"{key}_1", keep=True),
                       "Step 1 — the method set up in *Methods → Univariate calibration*: "
                       "compound, template, processing steps (left) and the laboratory "
                       "mixtures after processing with the measuring wavelength(s) marked "
                       "(right)."))
        d.tabs.setCurrentIndex(1)
        images.append((self.shot(d, f"{key}_2", keep=True),
                       "Step 2 — *Calibrate*: the calibration line of the standards, its "
                       "residuals and the regression statistics (slope, intercept, r, LOD, "
                       "LOQ)."))
        d.tabs.setCurrentIndex(2)
        images.append((self.shot(d, f"{key}_3"),
                       "Step 3 — *Determine*: found and taken concentrations and % recovery "
                       "of each mixture, with mean and RSD."))
        pr = d.predictions
        res = self.table([compound], pr["names"], [[f] for f in pr["found"]],
                         [[t] for t in pr["taken"]])
        settings = [f"Compound: {compound}; template: *{template}*"]
        name_of = {v: k for k, v in self.ids.items()}
        from spectro.core.operations import describe_step
        settings += ["Step: " + describe_step(st, lambda r: name_of.get(r, str(r)))
                     for st in d.method.steps]
        mp = d.method.measurement["params"]
        meas_txt = (uv.MEASUREMENTS[d.method.measurement["kind"]].label + " — " + ", ".join(
            f"{k} = {name_of.get(v, v) if k == 'reference' else fmt(v)}"
            for k, v in mp.items() if v is not None))
        settings.append("Measurement: " + meas_txt)
        settings.append(f"Calibration: slope {reg.slope:.5g}, intercept {reg.intercept:.3g}, "
                        f"r = {reg.r:.5f} (n = {reg.n})")
        if spike:
            res["headers"].append("Found − added")
            for row, f in zip(res["rows"], pr["found"]):
                row.append(f"{f - spike:.3f}")
        cal = self.p.spectra(d.cal_ids)
        calc = {"compound": compound, "signal": meas_txt, "direct": d.method.direct,
                "spike": spike,
                "cal": [(s.metadata["id"], s.concentrations[compound], d.method.signal(s, res_))
                        for s in cal],
                "test": [(i, t, sg) for i, t, sg in zip(pr["ids"], pr["taken"], pr["signals"])]}
        refs = self.refs_in(d.method.steps, d.method.measurement)
        data = {"sheets": {"Standards": list(d.cal_ids), "Divisor": refs,
                           ("Samples" if spike or mixtures is not None else "Mixtures"): list(mix)},
                "calc": calc}
        return {"images": images, "settings": settings + (notes or []), "data": data, **res}

    def refs_in(self, steps, measurement=None) -> list[int]:
        from spectro.core.operations import REGISTRY
        out = []
        for st in steps:
            for prm in REGISTRY[st["op"]].params:
                v = st["params"].get(prm.name)
                if prm.kind == "spectrum" and v is not None and int(v) not in out:
                    out.append(int(v))
        v = (measurement or {}).get("params", {}).get("reference")
        if v is not None and int(v) not in out:
            out.append(int(v))
        return out

    # ------------------------------------------------------------ finders
    def zero_crossing(self, c, lo, hi, steps):
        from spectro.core.operations import apply_pipeline
        s = apply_pipeline(self.unit(c), steps, lambda r: self.p.spectrum(r))
        return uv.zero_crossings(s, lo, hi)

    def equal(self, c, lam, lo, hi, steps=None):
        from spectro.core.operations import apply_pipeline
        s = apply_pipeline(self.unit(c), steps or [], lambda r: self.p.spectrum(r))
        return [w for w in uv.equal_amplitude_wavelengths(s, lam, lo, hi) if abs(w - lam) > 4]

    def best_lambda(self, c, steps, lo, hi, interferents=()):
        """λ with the best signal-to-noise for compound c after ``steps``,
        judged from its unit spectrum and simulated noise (as the method
        optimizer does), never from the test mixtures. Interferents must give
        ≈ 0 there."""
        from spectro.core.operations import apply_pipeline
        res = self.p.resolver()
        sig = apply_pipeline(self.unit(c), steps, res)
        rng = np.random.default_rng(42)
        noise = np.array([apply_pipeline(Spectrum(GRID, rng.normal(0, NOISE, GRID.size)),
                                         steps, res).values for _ in range(30)])
        x = sig.wavelengths
        ok = (x >= lo) & (x <= hi)
        sd = noise.std(axis=0)
        snr = np.divide(np.abs(sig.values), sd, out=np.zeros_like(sd), where=sd > 0)
        inter = np.zeros_like(sd)
        for k in interferents:
            inter += np.abs(apply_pipeline(self.unit(k), steps, res).values)
        if interferents:     # signal over interference + noise
            snr = np.divide(np.abs(sig.values), inter + sd, out=np.zeros_like(sd),
                            where=(inter + sd) > 0)
        snr[~ok] = 0
        return float(round(x[int(np.argmax(snr))], 1))

    def iso(self, a, b, lo, hi):
        return [w for w in uv.isoabsorptive_points(self.unit(a), self.unit(b)) if lo < w < hi]

    # ------------------------------------------------------------ primer
    def primer_shots(self, workbook: Path, workdir: Path) -> list[tuple[str, str]]:
        """Import a sample workbook into a new project, as a beginner would."""
        from spectro.storage.project import Project
        from spectro.ui.dialogs_data import ImportDialog
        p = Project.create(workdir / "beginner.spectro", "My first project")
        for c in ("A", "B"):
            p.add_compound(c)
        p.add_trial("Ratio difference example")
        self.win._attach(p)
        d = ImportDialog(self.win)
        d.resize(1250, 760)
        d.files = [str(workbook)]
        d._reparse()
        d._parse_names()
        d.table.selectRow(6)
        shots = [(self.shot(d, "primer_import", 760, keep=True),
                  "Importing a sample workbook: *File → Import spectra… → Add files…*, then "
                  "*Fill concentrations from names*. Each row is one spectrum; the Role column "
                  "comes from the sheet name and the A / B columns from the column titles.")]
        d._import()
        recs = p.records()
        from tools_select import select_ids
        select_ids(self.win, [r.id for r in recs if r.role == "standard"])
        shots.append((self.shot(self.win, "primer_project", 860),
                      "After import: the spectra are in the project tree with their roles "
                      "and concentrations; selecting them plots them. Every import is "
                      "recorded in the audit trail."))
        self.win.close_project()
        return shots

    # ------------------------------------------------------------ concept figures
    COLOURS = ["#2a78d6", "#d6862a", "#3a9a5b"]

    def _fig(self, n=2):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig, axes = plt.subplots(1, n, figsize=(12.5 if n == 2 else 7.5, 3.9))
        axes = list(np.atleast_1d(axes))
        for ax in axes:
            ax.grid(alpha=0.2)
            ax.set_xlabel("Wavelength (nm)")
        return fig, axes

    def _save_fig(self, fig, key, caption):
        import matplotlib.pyplot as plt
        fig.tight_layout()
        name = f"img/{key}_concept.png"
        fig.savefig(self.out / name, dpi=96)
        plt.close(fig)
        compact_png(self.out / name)
        return (name, caption)

    def _system_mix(self, compound):
        sysd = SYSTEMS[SYSTEM_OF[compound]]
        return sysd, sysd["mixtures"][0]

    def _zero_panel(self, ax, compound, marks=(), spans=()):
        sysd, mixc = self._system_mix(compound)
        total = np.zeros_like(GRID)
        for col, (c, u) in zip(self.COLOURS, sysd["pure"].items()):
            ax.plot(GRID, u * mixc[c], color=col, lw=1.8, label=f"{c} {mixc[c]:g} µg/mL alone")
            total += u * mixc[c]
        ax.plot(GRID, total, color="#222", lw=1.4, ls="--", label="the mixture (sum)")
        self._marks(ax, marks, spans)
        ax.set_title("Zero-order spectra: the mixture is the sum of its components", fontsize=10)
        ax.set_ylabel("Absorbance")
        ax.set_xlim(200, 400)
        ax.legend(frameon=False, fontsize=8)

    @staticmethod
    def _marks(ax, marks=(), spans=()):
        for lab, x in marks:
            ax.axvline(x, color="#555", lw=1.1)
            ax.annotate(lab, (x, 1), xycoords=("data", "axes fraction"), xytext=(3, -12),
                        textcoords="offset points", fontsize=8, color="#333")
        for lab, (a, b) in spans:
            ax.axvspan(a, b, color="#f2c14e", alpha=0.25)
            ax.annotate(lab, ((a + b) / 2, 0.02), xycoords=("data", "axes fraction"),
                        ha="center", fontsize=8, color="#7a5a00")

    def concept_univariate(self, key, compound, method):
        """Left: zero-order spectra; right: each compound processed ALONE by
        the method's steps — the analyte keeps a signal at the measuring λ, the
        interferents give ≈ 0 there (or a constant that the measurement removes)."""
        from spectro.core.operations import apply_pipeline
        res = self.p.resolver()
        sysd, mixc = self._system_mix(compound)
        fig, (a1, a2) = self._fig(2)
        p = method.measurement["params"]
        kind = method.measurement["kind"]
        if kind in ("amplitude",):
            marks, spans = [("λ", p["w1"])], []
        elif kind in ("difference", "weighted_difference", "ratio"):
            marks, spans = [("λ1", p["w1"]), ("λ2", p["w2"])], []
        else:
            marks, spans = [], [(kind, (p["w1"], p["w2"]))]
        for st in method.steps:
            q = st["params"]
            if "start" in q and "end" in q and st["op"] != "mean_center" and \
                    st["op"] != "divide_centered_ratio":
                spans.append(("plateau", (q["start"], q["end"])))
            if q.get("wavelength_end"):
                spans.append(("factor range", (q["wavelength"], q["wavelength_end"])))
            elif "wavelength" in q and st["op"] in ("factorized_recovery",
                                                     "spectrum_subtraction"):
                marks.append(("factor λ", q["wavelength"]))
        self._zero_panel(a1, compound)
        # ratio spectra are meaningless where the divisor is ≈ 0: hide that region
        keep = None
        for st in method.steps:
            ref = st["params"].get("reference") if st["op"] in (
                "divide", "divide_sum", "divide_centered_ratio") else \
                st["params"].get("divisor") if st["op"] in (
                    "ratio_subtraction", "constant_multiplication",
                    "extended_ratio_subtraction", "constant_center") else None
            if ref is not None:
                dv = res(ref)
                big = np.abs(dv.values) >= 0.05 * np.max(np.abs(dv.values))

                def keep(x, dv=dv, big=big):
                    return np.interp(x, dv.wavelengths, big.astype(float)) > 0.5
                break
        total = None
        for col, (c, u) in zip(self.COLOURS, sysd["pure"].items()):
            out = apply_pipeline(Spectrum(GRID, u * mixc[c]), method.steps, res)
            y = out.values.copy()
            if keep is not None:
                y[~keep(out.wavelengths)] = np.nan
            a2.plot(out.wavelengths, y, color=col, lw=2.2 if c == compound else 1.6,
                    label=f"{c} alone, processed" + ("  ← analyte" if c == compound else ""))
            total = y if total is None else total + y
        a2.plot(out.wavelengths, total, color="#222", lw=1.2, ls="--",
                label="mixture, processed")
        self._marks(a2, marks, spans)
        a2.axhline(0, color="#999", lw=0.8)
        a2.set_title("After the processing steps (each compound processed alone)",
                     fontsize=10)
        a2.set_ylabel("Processed signal")
        a2.legend(frameon=False, fontsize=8)
        others = [c for c in sysd["pure"] if c != compound]
        return self._save_fig(fig, key, (
            "Why it works — left: the spectra of the components and of the mixture "
            f"(they overlap). Right: each component processed alone with the method's steps. "
            f"At the marked wavelength(s) the processed {', '.join(others)} "
            "contribute nothing (≈ 0, or a constant that the measurement cancels), so the "
            f"mixture's signal there belongs to {compound} only."))

    def concept_zero(self, key, compound, marks=(), spans=(), caption=""):
        fig, (ax,) = self._fig(1)
        self._zero_panel(ax, compound, marks, spans)
        return self._save_fig(fig, key, caption)

    def concept_ratio(self, key, compound, divisor_id, marks=(), spans=(), caption="",
                      label="divisor"):
        sysd, mixc = self._system_mix(compound)
        div = self.p.spectrum(divisor_id).values
        ok = np.abs(div) >= 0.05 * np.max(np.abs(div))
        fig, (a1, a2) = self._fig(2)
        self._zero_panel(a1, compound)
        tot = np.zeros(ok.sum())
        for col, (c, u) in zip(self.COLOURS, sysd["pure"].items()):
            r = u[ok] * mixc[c] / div[ok]
            tot += r
            a2.plot(GRID[ok], r, ".", ms=2.2, color=col, label=f"{c} alone ÷ {label}")
        a2.plot(GRID[ok], tot, ".", ms=1.6, color="#222", label="mixture ÷ " + label)
        self._marks(a2, marks, spans)
        a2.set_title(f"Ratio spectra (÷ {label}); the divisor compound is a flat constant",
                     fontsize=10)
        a2.set_ylabel("Ratio amplitude")
        a2.legend(frameon=False, fontsize=8, markerscale=4)
        return self._save_fig(fig, key, caption)

    # ------------------------------------------------------------ finder screenshots
    def finder_shot(self, key, a_id, b_id, action, lo, hi, caption):
        from spectro.ui.dialogs_tools import FinderDialog
        self.win.selected_ids = lambda: [a_id, b_id]
        d = FinderDialog(self.win)
        del self.win.selected_ids
        d.resize(1300, 760)
        d.lo.setValue(lo)
        d.hi.setValue(hi)
        getattr(d, action)()
        return (self.shot(d, f"{key}_finder", 760),
                "How the wavelength was found — *Tools → Spectral finder*: " + caption)

    # ------------------------------------------------------------ examples
    D1 = {"order": 1, "delta_lambda": 4.0, "scaling": 10.0}

    def direct(self):
        return self.univariate("direct", "B", "Direct (zero order, λmax)", meas={"w1": 335.0},
                               notes=["B absorbs alone above 300 nm, so 335 nm (its second "
                                      "band) is free from A."])

    def dw(self):
        w2 = round(self.equal("B", 245.0, 205, 300)[0], 1)
        out = self.univariate("dw", "A", "Dual wavelength", meas={"w1": 245.0, "w2": w2},
                              notes=[f"Finder: B has the same absorbance at 245 nm and "
                                     f"{w2} nm."])
        out["images"].insert(1, self.finder_shot(
            "dw", self.ids["B′"], self.ids["B′"], "_equal", 245, 400,
            "spectrum A = the B′ standard, *From* = 245 nm (λmax of A) → *λ where A equals its "
            f"value at From*: B has the same absorbance at {w2} nm."))
        return out

    def idw(self):
        return self.univariate("idw", "A", "Induced dual wavelength",
                               meas={"w1": 245.0, "w2": 230.0, "reference": self.ids["B′"]},
                               notes=["F = A_B(245)/A_B(230) is computed from the B′ standard."])

    def auc1(self):
        return self.univariate("auc1", "B", "Area under curve (single component)",
                               meas={"w1": 320.0, "w2": 350.0, "baseline": "zero"})

    def _special(self, key, tab, setup, run, concept=None, divisor=None):
        from spectro.ui.dialogs_methods import SpecialDialog
        d = SpecialDialog(self.win)
        d.resize(1300, 820)
        d.tabs.setCurrentIndex(tab)
        setup(d)
        getattr(d, run)()
        img = self.shot(d, key)
        _, data = d.last
        images = ([concept] if concept else []) + [(img, (
            f"The method in *Methods → Binary two-signal methods → {d.tabs.tabText(tab)}*: "
            "settings (left), the standards and mixtures used (checked), and the results with "
            "recoveries (bottom)."))]
        wb = {"sheets": {"Standards": sorted(self._std_ids(data["X"], data["Y"])),
                         "Divisor": [divisor] if divisor else [], "Mixtures": data["ids"]}}
        names = [self.p.record(i).name for i in data["ids"]]
        found = [[r["X"], r["Y"]] for r in data["found"]]
        taken = [[self.p.record(i).concentrations.get(data["X"]),
                  self.p.record(i).concentrations.get(data["Y"])] for i in data["ids"]]
        res = self.table([data["X"], data["Y"]], names, found, taken)
        return {"images": images, "data": wb,
                "settings": [f"{k}: {fmt(v) if isinstance(v, float) else v}"
                             for k, v in data["params"].items()
                             if not isinstance(v, (list, tuple, dict)) and k != "divisor"],
                **res}

    def _std_ids(self, *comps):
        return {self.ids[f"{c} {v}"] for c in comps for v in SYSTEMS[SYSTEM_OF[c]]["levels"]}

    def q(self):
        iso = self.iso("A", "B", 230, 300)[0]

        def setup(d):
            d.q_x.setCurrentText("A")
            d.q_y.setCurrentText("B")
            d.q_iso.setValue(iso)
            d.q_w2.setValue(245.0)
            check_only(d.q_xs, self._std_ids("A"))
            check_only(d.q_ys, self._std_ids("B"))
            check_only(d.q_mix, self.mix["Binary A + B"])
        concept = self.concept_zero(
            "q", "A", [("iso", iso), ("λ2", 245.0)],
            caption=f"Why it works — at the isoabsorptive point ({iso:.1f} nm) A and B of equal "
                    "concentration absorb the same, so the mixture's absorbance there measures "
                    "A + B; the ratio of the mixture's absorbances at λ2 and iso tells how the "
                    "total is split.")
        out = self._special("q", 0, setup, "_run_q", concept)
        out["images"].insert(1, self.finder_shot(
            "q", self.ids["A 12"], self.ids["B 12"], "_iso", 220, 300,
            "spectrum A = an A standard, spectrum B = a B standard, *Divide by total "
            "concentration* ticked → *Isoabsorptive points of A and B*."))
        out["settings"].insert(0, f"Isoabsorptive point (finder): {iso:.1f} nm; λ2 = 245 nm "
                                  "(λmax of A)")
        return out

    def as_(self):
        iso = self.iso("A", "B", 230, 300)[0]

        def setup(d):
            d.as_x.setCurrentText("B")
            d.as_y.setCurrentText("A")
            d.as_iso.setValue(iso)
            d.as_w2.setValue(335.0)
            check_only(d.as_xs, self._std_ids("B"))
            check_only(d.as_all, self._std_ids("A", "B"))
            check_only(d.as_mix, self.mix["Binary A + B"])
        concept = self.concept_zero(
            "as", "A", [("iso", iso), ("λ2", 335.0)],
            caption="Why it works — B absorbs alone at λ2 (335 nm): its reading there, times "
                    "B's own ratio A(iso)/A(λ2), gives B's absorbance at the isoabsorptive point; "
                    "the rest of the absorbance at iso belongs to A.")
        out = self._special("as", 1, setup, "_run_as", concept)
        out["settings"].insert(0, f"X = B (absorbs alone at λ2 = 335 nm); iso = {iso:.1f} nm")
        return out

    def aas(self):
        iso = self.iso("A", "B", 230, 300)[0]
        w1 = self.equal("B", iso, 205, 300)[0]

        def setup(d):
            d.aas_x.setCurrentText("A")
            d.aas_y.setCurrentText("B")
            d.aas_w2.setValue(iso)
            d.aas_w1.setValue(round(w1, 1))
            check_only(d.aas_xs, self._std_ids("A"))
            check_only(d.aas_ys, self._std_ids("B"))
            check_only(d.aas_mix, self.mix["Binary A + B"])
        concept = self.concept_zero(
            "aas", "A", [("λ2 = iso", iso), ("λ1", w1)],
            caption="Why it works — B absorbs the same at λ1 and λ2 (= isoabsorptive point), so "
                    "A(λ2) − A(λ1) of the mixture depends on A only; A's share at λ2 follows and "
                    "the remainder is B.")
        out = self._special("aas", 2, setup, "_run_aas", concept)
        out["settings"].insert(0, f"λ2 = isoabsorptive point {iso:.1f} nm; λ1 = {w1:.1f} nm "
                                  "where B has the same absorbance")
        return out

    def am(self):
        iso = self.iso("A", "B", 230, 300)[0]
        unit = self.ids["B unit"]

        def setup(d):
            d.am_x.setCurrentText("A")
            d.am_y.setCurrentText("B")
            d.am_iso.setValue(iso)
            d.am_p1.setValue(310.0)
            d.am_p2.setValue(350.0)
            d.am_div.setCurrentIndex(d.am_div.findData(unit))
            d.am_divc.setValue(1.0)
            check_only(d.am_std, self._std_ids("A", "B"))
            check_only(d.am_mix, self.mix["Binary A + B"])
        concept = self.concept_ratio(
            "am", "A", unit, [("iso", iso)], [("plateau → B", (310.0, 350.0))],
            caption="Why it works — divided by B at 1 µg/mL, B becomes a flat line equal to its "
                    "concentration; at the isoabsorptive point the ratio equals A + B. Plateau → "
                    "B, iso − plateau → A.", label="B (1 µg/mL)")
        out = self._special("am", 3, setup, "_run_am", concept, unit)
        out["settings"].insert(0, f"Divisor: unit-concentration B (Normalize → concentration); "
                                  f"iso {iso:.1f} nm; plateau 310–350 nm")
        return out

    def iam(self):
        def setup(d):
            d.iam_x.setCurrentText("A")
            d.iam_y.setCurrentText("B")
            d.iam_w.setValue(245.0)
            d.iam_p1.setValue(310.0)
            d.iam_p2.setValue(350.0)
            check_only(d.iam_xs, self._std_ids("A"))
            check_only(d.iam_ys, self._std_ids("B"))
            check_only(d.iam_mix, self.mix["Binary A + B"])
        concept = self.concept_ratio(
            "iam", "A", self.ids["B unit"], [("λ for A", 245.0)],
            [("plateau → B", (310.0, 350.0))],
            caption="Why it works — divided by B at 1 µg/mL, B is a flat line equal to C_B "
                    "(plateau); at 245 nm the ratio is C_B + r·C_A, with r = a_A/a_B from the A "
                    "standards.", label="B (1 µg/mL)")
        return self._special("iam", 4, setup, "_run_iam", concept)

    def hpsam(self):
        tid = self.trial["Binary A + B"]
        base = {"A": 6.0, "B": 12.0}
        pure = SYSTEMS["Binary A + B"]["pure"]
        ids = []
        for i, add in enumerate((0, 4, 8, 12)):
            y = pure["A"] * (base["A"] + add) + pure["B"] * base["B"]
            y = y + np.random.default_rng(500 + i).normal(0, NOISE, GRID.size)
            ids.append(self.p.add_spectrum(Spectrum(GRID, y, name=f"Sample + {add} A",
                                                    concentrations={"A": float(add)}),
                                           tid, role="sample"))
        self.spiked = ids
        w2 = self.equal("B", 245.0, 205, 300)[0]
        from spectro.ui.dialogs_methods import SpecialDialog
        d = SpecialDialog(self.win)
        d.tabs.setCurrentIndex(5)
        d.h_x.setCurrentText("A")
        d.h_w1.setValue(245.0)
        d.h_w2.setValue(round(w2, 1))
        check_only(d.h_set, ids)
        d._run_h()
        img = self.shot(d, "hpsam")
        r = d.last[1]
        fig, (a1, a2) = self._fig(2)
        added = [0, 4, 8, 12]
        for col, sid, add in zip(["#2a78d6", "#3a9a5b", "#d6862a", "#9b59b6"], ids, added):
            sp = self.p.spectrum(sid)
            a1.plot(sp.wavelengths, sp.values, color=col, lw=1.5, label=f"sample + {add} A")
        self._marks(a1, [("λ1", 245.0), ("λ2", w2)])
        a1.set_title("Sample with standard additions of A", fontsize=10)
        a1.set_ylabel("Absorbance")
        a1.legend(frameon=False, fontsize=8)
        a1.set_xlim(200, 320)
        xs = np.array([-r["X"] - 1, 13.0])
        for lab, ln, col in (("at λ1", r["line1"], "#2a78d6"), ("at λ2", r["line2"], "#d6862a")):
            sps = self.p.spectra(ids)
            lam = 245.0 if lab == "at λ1" else w2
            a2.plot(added, [sp.value_at(lam) for sp in sps], "o", color=col)
            a2.plot(xs, ln["slope"] * xs + ln["intercept"], color=col, label=lab)
        a2.plot([-r["X"]], [r["A_H"]], "s", color="#222", ms=8, label="H point")
        a2.axvline(0, color="#999", lw=0.8)
        a2.set_xlabel("A added (µg/mL)")
        a2.set_ylabel("Absorbance")
        a2.set_title("The two addition lines cross at H(−C_A, A_B)", fontsize=10)
        a2.legend(frameon=False, fontsize=8)
        images = [self._save_fig(fig, "hpsam", (
            "Why it works — B absorbs the same at λ1 and λ2, so it shifts both lines by the "
            "same amount; the lines therefore cross at the x-value −C_A, whatever B's "
            "concentration.")),
            (img, "*Binary two-signal methods → Standard addition: H-point*: analyte, λ1, λ2 "
                  "and the sample with its additions (A concentration = amount added).")]
        return {"images": images,
                "data": {"sheets": {"Samples": ids},
                         "note": "In the Samples sheet the A value of each spectrum is the "
                                 "amount of A ADDED (0 = the sample alone); the sample itself "
                                 "contains A 6 and B 12 µg/mL."},
                "settings": ["Sample: A 6 µg/mL + B 12 µg/mL; additions of A: 0, 4, 8, 12 µg/mL",
                             f"λ1 = 245 nm, λ2 = {w2:.1f} nm (B equal, finder)"],
                "headers": ["Quantity", "Value"],
                "rows": [["A found (−C_H)", f"{r['X']:.3f} µg/mL (true 6)"],
                         ["Recovery", f"{100 * r['X'] / 6:.2f} %"],
                         ["A_H (interferent B signal)", f"{r['A_H']:.4f}"]],
                "summary": []}

    def _progressive(self, key, tab, setup, run):
        from spectro.ui.dialogs_methods import ProgressiveDialog
        d = ProgressiveDialog(self.win)
        d.resize(1300, 1050)
        d.tabs.setCurrentIndex(tab)
        setup(d)
        getattr(d, run)()
        img = self.shot(d, key, 1050)
        _, data = d.last
        comps = data["compounds"]
        names = [self.p.record(i).name for i in data["ids"]]
        taken = [[self.p.record(i).concentrations.get(c) for c in comps] for i in data["ids"]]
        res = self.table(comps, names, data["found"], taken)
        m = d.fitted[0]
        if isinstance(m, uv.AmplitudeCentering):
            marks = [("λc", m.wavelength)] + [(f"λ2 {c}", df["w2"])
                                               for c, df in m.differences.items()]
            spans = [(f"plateau → {m.divisor_compound}", m.plateau)] if m.plateau else []
            concept = self.concept_ratio(
                key, comps[0], m.divisor, marks, spans, label=f"{m.divisor_compound}′",
                caption="Why it works — in the ratio spectrum the divisor compound "
                        f"({m.divisor_compound}) is a flat constant. Every compound is read at "
                        "the same λc: from the plateau, from an amplitude difference at a pair "
                        "where the others are equal (they cancel), or as what is left after "
                        "subtracting the others.")
            divisor = [m.divisor]
        else:
            concept = self.concept_zero(
                key, comps[0], [(f"{c}", w) for c, w in m.order],
                caption="Why it works — at each marked wavelength the compound named there is "
                        "added to those already resolved (first: absorbs alone; second: "
                        "overlapped only by the first; …); their known contributions are "
                        "subtracted step by step.")
            divisor = []
        images = [concept, (img, f"*Methods → Progressive resolution → {d.tabs.tabText(tab)}*: "
                                 "the settings, the standards and mixtures used, and the "
                                 "results with recoveries. *Save method* stores it for reuse.")]
        wb = {"sheets": {"Standards": sorted(self._std_ids(*comps)), "Divisor": divisor,
                         "Mixtures": data["ids"]}}
        return {"images": images, "data": wb, "settings": [m.describe()], **res}

    def _ac(self, d, comps, divisor, divc, lam, plateau=None, rows=(), sub=None, unified=False):
        for i in range(d.ac_comps.count()):
            it = d.ac_comps.item(i)
            it.setCheckState(Qt.Checked if it.text() in comps else Qt.Unchecked)
        d.ac_div.setCurrentIndex(d.ac_div.findData(self.ids[divisor]))
        d.ac_divc.setCurrentText(divc)
        d.ac_w.setValue(lam)
        d.ac_plateau.setChecked(plateau is not None)
        if plateau:
            d.ac_p1.setValue(plateau[0])
            d.ac_p2.setValue(plateau[1])
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                d.ac_diff.setItem(i, j, QTableWidgetItem("" if v is None else str(v)))
        d.ac_sub.setCurrentText(sub or "(none)")
        d.ac_unified.setChecked(unified)
        sys_name = SYSTEM_OF[comps[0]]
        check_only(d.ac_std, self._std_ids(*comps))
        check_only(d.ac_mix, self.mix[sys_name])

    def _ratio_unit(self, c, div):
        u = self.unit(c)
        dv = np.interp(GRID, GRID, self.p.spectrum(self.ids[div]).values)
        return u.with_values(u.values / dv)

    def _equal_ratio(self, c, div, lam, lo, hi):
        r = self._ratio_unit(c, div)
        return [w for w in uv.equal_amplitude_wavelengths(r, lam, lo, hi) if abs(w - lam) > 4]

    def cvad(self):
        return self._progressive("cvad", 0, lambda d: self._ac(
            d, ["A", "B"], "B′", "B", 251.0, rows=[("A", 261.0, 251.0, None)], sub="B"),
            "_run_ac")

    def aac(self):
        return self._progressive("aac", 0, lambda d: self._ac(
            d, ["X", "Y", "Z"], "Z′", "Z", 275.0, plateau=(340.0, 380.0),
            rows=[("X", 275.0, 240.0, "Y")], sub="Y"), "_run_ac")

    def optimizer_config(self, divisor_compound, plateau: bool):
        """Best amplitude-centering settings proposed by the method optimizer
        (screened on the pure standards with simulated noise and ±0.5 nm)."""
        from spectro.core.optimizer import OptimizerInput, _ProgressiveScreen, _Screen
        comps = ["X", "Y", "Z"]
        std = {c: [s for s in self.p.spectra(sorted(self._std_ids(c)))] for c in comps}
        divs = {c: self.p.spectrum(self.ids[f"{c}′"]) for c in comps}
        inp = OptimizerInput(comps, std, divs, families={"amplitude_centering"})
        for c in comps:
            for s_ in inp.standards[c]:
                s_.concentrations = {**{k: 0.0 for k in comps}, **s_.concentrations}
        ps = _ProgressiveScreen(_Screen(inp))
        best = None
        for m, z in ps.amplitude_centering():
            if (divisor_compound and z != divisor_compound) or bool(m.plateau) != plateau:
                continue
            ev = ps.evaluate(m, ps.sc.resolve("div:" + z))
            if ev and (best is None or max(e["robust"] for e in ev.values()) < best[0]):
                best = (max(e["robust"] for e in ev.values()), m, z)
        return best[1], best[2]

    def macm(self):
        m, z = self.optimizer_config(None, plateau=False)
        lam = round(m.wavelength, 1)
        rows = [(c, lam, round(d["w2"], 1), None) for c, d in m.differences.items()]
        out = self._progressive("macm", 0, lambda d: self._ac(
            d, ["X", "Y", "Z"], f"{z}′", z, lam, rows=rows, sub=z), "_run_ac")
        out["settings"].append("Settings proposed by Tools → Method optimizer (λc and the "
                               "partner wavelengths where the other compound's ratio "
                               "amplitudes are equal).")
        return out

    def ridss(self):
        lam = self.iso("X", "Y", 230, 300)[0]
        wx = min(self._equal_ratio("X", "Z′", lam, 210, 320), key=lambda w: abs(w - lam))
        out = self._progressive("ridss", 0, lambda d: self._ac(
            d, ["X", "Y", "Z"], "Z′", "Z", round(lam, 1), plateau=(340.0, 380.0),
            rows=[("Y", round(lam, 1), round(wx, 1), None)], sub="X", unified=True), "_run_ac")
        out["settings"].append(f"Isoabsorptive point of X and Y: {lam:.1f} nm; X/Z′ equal at "
                               f"{wx:.1f} nm")
        return out

    def mafm(self):
        def setup(d):
            for i, (c, lam, q) in enumerate([("U", "335", "300"), ("V", "280", ""),
                                             ("W", "246", "")]):
                for j, v in enumerate((c, lam, q)):
                    d.af_table.setItem(i, j, QTableWidgetItem(v))
            check_only(d.af_std, self._std_ids("U", "V", "W"))
            check_only(d.af_mix, self.mix["Ternary U + V + W"])
        return self._progressive("mafm", 1, setup, "_run_af")

    def _equations(self, key, comps, signals, kaiser=False, steps=None, intercept=False):
        from spectro.ui.dialogs_methods import EquationsDialog
        from spectro.ui.widgets import fill_table
        d = EquationsDialog(self.win)
        d.resize(1400, 820)
        for i in range(d.comps.count()):
            it = d.comps.item(i)
            it.setCheckState(Qt.Checked if it.text() in comps else Qt.Unchecked)
        check_only(d.cal, self._std_ids(*comps))
        check_only(d.test, self.mix[SYSTEM_OF[comps[0]]])
        if steps:
            d.pipe.set_steps(steps)
        if kaiser:
            d._kaiser()
        else:
            fill_table(d.signals, ["Kind (amplitude/area)", "λ1 (nm)", "λ2 (nm, area only)"],
                       signals, editable=True)
            d.intercept.setChecked(intercept)
        d._fit()
        d._predict()
        img = self.shot(d, key)
        marks, spans = [], []
        for sg in d.model.signals:
            q = sg["params"]
            if sg["kind"] == "area":
                spans.append(("area", (q["w1"], q["w2"])))
            else:
                marks.append(("λ", q["w1"]))
        concept = self.concept_zero(
            key, comps[0], marks, spans,
            caption="Why it works — each marked signal of the mixture is the sum of every "
                    "compound's contribution (absorptivity × concentration); with one signal per "
                    "compound (or more) the system of equations is solved for all "
                    "concentrations at once.")
        pr = d.predictions
        names = [self.p.record(i).name for i in pr["ids"]]
        taken = [[self.p.record(i).concentrations.get(c) for c in pr["compounds"]]
                 for i in pr["ids"]]
        res = self.table(pr["compounds"], names, pr["found"], taken)
        sig = [f"{s['kind']} {s['params']['w1']:g}" + (f"–{s['params']['w2']:g}"
                                                        if s["kind"] == "area" else "") + " nm"
               for s in d.model.signals]
        return {"images": [concept, (img, "*Methods → Equation methods*: compounds, signals, "
                                          "the absorptivity matrix K after *Fit* (top right) "
                                          "and the results after *Determine*.")],
                "data": {"sheets": {"Standards": sorted(self._std_ids(*comps)),
                                    "Mixtures": pr["ids"]}},
                "settings": ["Signals: " + "; ".join(sig), d.summary.text().split("\n")[0]],
                **res}

    def vierordt(self):
        return self._equations("vierordt", ["A", "B"], [["amplitude", "245", ""],
                                                         ["amplitude", "268", ""]])

    def bivariate(self):
        return self._equations("bivariate", ["A", "B"], None, kaiser=True)

    def auc_eq(self):
        return self._equations("auc_eq", ["X", "Y", "Z"],
                               [["area", "230", "245"], ["area", "285", "295"],
                                ["area", "340", "370"]])

    def zc(self):
        zcs = self.zero_crossing("B", 230, 260, [{"op": "derivative", "params": self.D1}])
        lam = round(zcs[0], 1)
        out = self.univariate("zc", "A", "Zero-crossing derivative (D1)",
                              params={0: {"scaling": 10.0}}, meas={"w1": lam},
                              notes=[f"Finder: B's D1 crosses zero at {lam} nm."])
        out["images"].insert(1, self.finder_shot(
            "zc", self.ids["B′ D1"], self.ids["B′ D1"], "_zero", 230, 260,
            "spectrum A = the D1 spectrum of B′ (made once with *Process → Derivative*), "
            f"range 230–260 nm → *Zero-crossing points of A*: {lam} nm."))
        return out

    def p2p(self):
        return self.univariate("p2p", "B", "Derivative peak-to-peak", params={0: {"scaling": 10.0}},
                               meas={"w1": 310.0, "w2": 365.0})

    def d1dwl(self):
        from spectro.core.operations import apply_pipeline
        steps = [{"op": "derivative", "params": self.D1}]
        da = apply_pipeline(self.unit("A"), steps)
        best = None
        for w1 in np.arange(225.0, 300.0, 1.0):
            for w2 in self.equal("B", w1, 205, 330, steps):
                gain = abs(da.value_at(w1) - da.value_at(w2))
                if best is None or gain > best[0]:
                    best = (gain, float(w1), round(w2, 1))
        _, w1, w2 = best
        return self.univariate("d1dwl", "A", "Dual wavelength in derivative mode",
                               params={0: {"scaling": 10.0}}, meas={"w1": w1, "w2": w2},
                               notes=[f"Finder: B's D1 is equal at {w1:g} and {w2} nm; of all "
                                      "such pairs this one gives A the largest difference."])

    def dd1(self):
        return self.univariate("dd1", "A", "Derivative ratio (DD1)",
                               refs={0: {"reference": "B′"}}, params={1: {"scaling": 10.0}},
                               meas={"w1": 254.0})

    def _template_steps(self, template, refs):
        steps = [dict(st, params=dict(st["params"])) for st in uv.TEMPLATES[template]["steps"]]
        for i, kv in refs.items():
            for k, v in kv.items():
                steps[i]["params"][k] = self.ids[v]
        return steps

    def d1dr(self):
        t = "Derivative ratio of D1 spectra (D1 DR)"
        lam = self.best_lambda("A", self._template_steps(t, {1: {"reference": "B′"}}), 215, 300)
        return self.univariate("d1dr", "A", t, refs={1: {"reference": "B′"}}, meas={"w1": lam},
                               notes=[f"{lam} nm: best signal-to-noise of the processed A "
                                      "spectrum (spikes where B's D1 crosses zero avoided)."])

    def mcr(self):
        return self.univariate("mcr", "A", "Mean centering of ratio spectra (MCR)",
                               refs={0: {"reference": "B′"}},
                               params={1: {"start": 220.0, "end": 300.0}}, meas={"w1": 240.0})

    def mcr3(self):
        t = "Mean centering of ratio spectra (ternary)"
        refs = {0: {"reference": "Z′"}, 2: {"reference": "Y′", "divisor": "Z′"}}
        rng = {1: {"start": 220.0, "end": 300.0}, 2: {"start": 220.0, "end": 300.0},
               3: {"start": 220.0, "end": 300.0}}
        steps = self._template_steps(t, refs)
        for i, kv in rng.items():
            steps[i]["params"].update(kv)
        lam = self.best_lambda("X", steps, 222, 298)
        return self.univariate("mcr3", "X", t, refs=refs, params=rng, meas={"w1": lam},
                               notes=[f"{lam} nm: best signal-to-noise of the processed X "
                                      "spectrum."])

    def sdr(self):
        t = "Successive derivative ratio (ternary)"
        refs = {0: {"reference": "Y′"}, 2: {"reference": "D1(Z′/Y′)"}}
        lam = self.best_lambda("X", self._template_steps(t, refs), 215, 320)
        return self.univariate("sdr", "X", t, refs=refs, meas={"w1": lam},
                               notes=["Second divisor: D1 of (Z′ ÷ Y′), made once with "
                                      "Process → pipeline and stored as a derived spectrum.",
                                      f"{lam} nm: best signal-to-noise of the processed X "
                                      "spectrum."])

    def dd(self):
        t = "Double divisor ratio derivative (ternary)"
        refs = {0: {"reference": "Y′", "reference2": "Z′"}}
        steps = self._template_steps(t, refs)
        steps[1]["params"]["scaling"] = 10.0
        lam = self.best_lambda("X", steps, 215, 330, interferents=("Y", "Z"))
        return self.univariate("dd", "X", t, refs=refs, params={1: {"scaling": 10.0}},
                               meas={"w1": lam},
                               notes=[f"{lam} nm: largest X signal relative to the residual Y "
                                      "and Z signals (from the pure spectra).",
                                      "Approximate: Y and Z cancel only in the proportion of the "
                                      "divisor (here 12 : 24). Compare the recoveries of mixtures "
                                      "with different Y : Z ratios."])

    def rd(self):
        return self.univariate("rd", "A", "Ratio difference (RD)", refs={0: {"reference": "B′"}},
                               meas={"w1": 245.0, "w2": 275.0})

    def dad(self):
        lam = 240.0
        w2 = min(self._equal_ratio("Y", "Z′", lam, 210, 320), key=lambda w: abs(w - lam))
        return self.univariate("dad", "X", "Dual amplitude difference (ternary)",
                               refs={0: {"reference": "Z′"}},
                               meas={"w1": lam, "w2": round(w2, 1), "reference": self.ids["Y′"]},
                               notes=[f"Finder: Y/Z′ equal at {lam:g} and {w2:.1f} nm."])

    def cv(self):
        out = self.univariate("cv", "B", "Ratio plateau / constant value (Y)",
                              refs={0: {"reference": "B′"}}, meas={"w1": 310.0, "w2": 350.0})
        out["images"].insert(1, self.finder_shot(
            "cv", self.ids["Mix ÷ B′"], self.ids["Mix ÷ B′"], "_plateau", 300, 360,
            "spectrum A = a mixture divided by B′ (*Process → Divide by spectrum*), range "
            "300–360 nm → *Plateau regions of A*: the flat region used for the constant."))
        return out

    def concval(self):
        return self.univariate("concval", "B",
                               "Concentration value (unit divisor plateau, no regression)",
                               refs={0: {"reference": "B unit"}},
                               meas={"w1": 310.0, "w2": 350.0}, direct=True,
                               notes=["The plateau value is reported directly as µg/mL."])

    def rs(self):
        return self.univariate("rs", "A", "Ratio subtraction / spectrum subtraction (X)",
                               refs={0: {"divisor": "B′"}},
                               params={0: {"start": 310.0, "end": 350.0}}, meas={"w1": 245.0})

    def cm(self):
        return self.univariate("cm", "B", "Constant multiplication / SS-CM (extended Y)",
                               refs={0: {"divisor": "B′"}},
                               params={0: {"start": 310.0, "end": 350.0}}, meas={"w1": 268.0})

    def ers(self):
        return self.univariate("ers", "B", "Extended ratio subtraction (Y)",
                               refs={0: {"divisor": "B′", "reference": "A 12"}},
                               params={0: {"start": 310.0, "end": 350.0}}, meas={"w1": 268.0})

    def srs(self):
        return self.univariate("srs", "W", "Successive ratio subtraction (ternary SRS, Z)",
                               refs={0: {"divisor": "U′"}, 1: {"divisor": "V′"}},
                               params={0: {"start": 330.0, "end": 345.0},
                                       1: {"start": 278.0, "end": 290.0}}, meas={"w1": 246.0})

    def sss(self):
        return self.univariate("sss", "W", "Successive spectrum subtraction (ternary)",
                               refs={0: {"reference": "U′"}, 1: {"divisor": "V′"}},
                               params={0: {"wavelength": 330.0, "wavelength_end": 345.0},
                                       1: {"start": 278.0, "end": 290.0}}, meas={"w1": 246.0})

    def ss(self):
        return self.univariate("ss", "A", "Spectrum subtraction",
                               refs={0: {"reference": "B′"}},
                               params={0: {"wavelength": 320.0, "wavelength_end": 350.0}},
                               meas={"w1": 245.0})

    def fzm(self):
        zcs = self.zero_crossing("B", 230, 260, [{"op": "derivative",
                                                  "params": {"order": 1, "delta_lambda": 4.0}}])
        lam = round(zcs[0], 1)
        return self.univariate("fzm", "A", "Factorized zero-order (FZM)",
                               refs={0: {"reference": "A 12"}}, params={0: {"wavelength": lam}},
                               meas={"w1": 245.0},
                               notes=[f"λ = {lam} nm, zero-crossing of B's D1 (finder)."])

    def cc(self):
        return self.univariate("cc", "A", "Constant center",
                               refs={0: {"divisor": "B′", "reference": "A 12"}},
                               params={0: {"w1": 240.0, "w2": 255.0, "target": "X"}},
                               meas={"w1": 245.0})

    def dt(self):
        return self.univariate("dt", "B", "Derivative transformation (DT, recover zero order)",
                               refs={0: {"reference": "B′"}},
                               params={0: {"wavelength": 305.0, "wavelength_end": 320.0}},
                               meas={"w1": 268.0},
                               notes=["A's D1 is zero above 300 nm, B's is not: 305–320 nm."])

    def ds(self):
        return self.univariate("ds", "B",
                               "Derivative subtraction–constant multiplication (DS-CM, Y in D1)",
                               refs={1: {"divisor": "B′"}},
                               params={1: {"start": 305.0, "end": 320.0}}, meas={"w1": 312.0},
                               notes=["DS-CM shown (B's D1 spectrum recovered); the DS template "
                                      "gives A's D1 spectrum the same way."])

    def chemo(self):
        from spectro.ui.dialogs_methods import ChemometricsDialog
        d = ChemometricsDialog(self.win)
        d.resize(1450, 900)
        tid = self.trial["Ternary X + Y + Z"]
        train = [r.id for r in self.p.records(trial_id=tid) if r.role == "calibration"]
        check_only(d.cal, train)
        check_only(d.test, self.mix["Ternary X + Y + Z"])
        for i in range(d.comps.count()):
            it = d.comps.item(i)
            it.setCheckState(Qt.Checked if it.text() in ("X", "Y", "Z") else Qt.Unchecked)
        d.ranges.setText("220-390")
        rows = []
        for model in ("CLS", "ILS", "PCR", "PLS2", "MCR-ALS", "ANN", "SVR"):
            d.mtype.setCurrentText(model)
            if model in ("PCR", "PLS2"):
                d._cv()
            if model == "ILS":
                d.wls.setText("240, 255, 275, 290, 355")
            d._fit()
            if model == "PLS2":
                d.tabs.setCurrentIndex(0)
                shots = [(self.shot(d, "chemo_1", 900, keep=True),
                          "Step 1 — *Cross-validate* (PLS2): RMSECV against the number of "
                          "components for each compound; the suggested number is set "
                          "automatically.")]
                d.tabs.setCurrentIndex(1)
                shots.append((self.shot(d, "chemo_2", 900, keep=True),
                              "Step 2 — *Fit + predict*: predicted vs actual concentrations of "
                              "the validation mixtures, with RMSEP and recoveries."))
                d.tabs.setCurrentIndex(3)
                shots.append((self.shot(d, "chemo_3", 900, keep=True),
                              "Step 3 — outlier check: Hotelling T² vs Q residuals with 95/99 % "
                              "limits; flagged spectra can be excluded with one click."))
            recs = {}
            for r in range(d.results.rowCount()):
                for j, c in enumerate(("X", "Y", "Z")):
                    it = d.results.item(r, 3 + 3 * j)
                    if it and it.text():
                        recs.setdefault(c, []).append(float(it.text()))
            rows.append([model + (f" ({d.ncomp.value()} components)"
                                  if model in ("PCR", "PLS2") else "")]
                        + [f"{describe(recs[c])['mean']:.2f} ± {describe(recs[c])['sd']:.2f}"
                           for c in ("X", "Y", "Z")])
        fig, (a1, a2) = self._fig(2)
        for sp in self.p.spectra(train):
            a1.plot(sp.wavelengths, sp.values, color="#8aa9d6", lw=0.8)
        a1.set_title("Training set: 25 mixtures of a calibration design", fontsize=10)
        a1.set_ylabel("Absorbance")
        cs = [self.p.record(i).concentrations for i in train]
        a2.scatter([c["X"] for c in cs], [c["Y"] for c in cs], c=[c["Z"] for c in cs],
                   cmap="viridis", s=40)
        a2.set_xlabel("X (µg/mL)")
        a2.set_ylabel("Y (µg/mL)")
        a2.set_title("Design: 5 levels per compound, uncorrelated (colour = Z)", fontsize=10)
        concept = self._save_fig(fig, "chemo", (
            "Why it works — the model learns from many mixtures whose composition varies "
            "independently for each compound how the WHOLE spectrum changes with each "
            "concentration; it then predicts all compounds of a new spectrum at once."))
        return {"images": [concept] + shots,
                "data": {"sheets": {"Calibration": train,
                                    "Mixtures": self.mix["Ternary X + Y + Z"]}},
                "settings": ["Training set: 25-mixture Brereton design (Tools → Calibration "
                             "design), 220–390 nm; validation: the 5 laboratory mixtures"],
                "headers": ["Model", "X recovery %", "Y recovery %", "Z recovery %"],
                "rows": rows, "summary": []}

    def stdadd(self):
        from spectro.ui.dialogs_methods import UnivariateDialog
        if not hasattr(self, "spiked"):
            self.hpsam()
        d = UnivariateDialog(self.win)
        d.compound.setCurrentText("A")
        d.template.setCurrentText("Ratio difference (RD)")
        st = d.pipe.get_steps()
        st[0]["params"]["reference"] = self.ids["B′"]
        d.pipe.set_steps(st)
        m = d.meas.get()
        m["params"].update(w1=245.0, w2=275.0)
        d.meas.set(m)
        d._calibrate()
        check_only(d.test, self.spiked)
        d._std_addition()
        d.tabs.setCurrentIndex(2)
        img = self.shot(d, "stdadd")
        sa = d.last[2]
        fig, (ax,) = self._fig(1)
        add = [r["added"] for r in sa["rows"]]
        found = [r["found"] for r in sa["rows"]]
        reg = sa["regression"]
        xs = np.array([-sa["extrapolated"] - 0.5, max(add) + 1])
        ax.plot(add, found, "o", color="#2a78d6")
        ax.plot(xs, reg["slope"] * xs + reg["intercept"], color="#2a78d6")
        ax.axvline(0, color="#999", lw=0.8)
        ax.axhline(0, color="#999", lw=0.8)
        ax.plot([-sa["extrapolated"]], [0], "s", color="#222")
        ax.set_xlabel("A added (µg/mL)")
        ax.set_ylabel("A found (µg/mL)")
        ax.set_title("Found vs added: slope ≈ 1 means 100 % recovery", fontsize=10)
        concept = self._save_fig(fig, "stdadd", (
            "Why it works — each addition should raise the found amount by exactly the amount "
            "added (slope 1); the line crosses zero at minus the amount already in the sample."))
        rows = [[r["name"], f"{r['added']:g}", f"{r['found']:.3f}",
                 "" if r["recovered"] is None else f"{r['recovered']:.3f}",
                 "" if r["recovery"] is None else f"{r['recovery']:.2f}"] for r in sa["rows"]]
        return {"images": [concept, (img, "*Standard addition…* in the Univariate dialog with a "
                                          "calibrated method: each spiked sample, the amount "
                                          "recovered and its recovery.")],
                "data": {"sheets": {"Standards": sorted(self._std_ids("A")),
                                    "Divisor": [self.ids["B′"]], "Samples": self.spiked},
                         "note": "In the Samples sheet the A value of each spectrum is the "
                                 "amount of A ADDED to the sample (0 = the sample alone)."},
                "settings": ["Method: ratio difference for A (÷ B′, 245 − 275 nm)",
                             f"A in the sample: {sa['sample']:.3f} µg/mL (true 6); "
                             f"extrapolated {sa['extrapolated']:.3f}"],
                "headers": ["Spectrum", "A added", "Found", "Recovered", "Recovery %"],
                "rows": rows,
                "summary": [f"Mean recovery of the additions {sa['mean_recovery']:.2f} %"]}

    def enrich(self):
        tid = self.trial["Binary A + B"]
        pure = SYSTEMS["Binary A + B"]["pure"]
        ids = []
        for i in range(3):
            y = pure["A"] * (20.0 + 0.0) + pure["B"] * (1.0 + 5.0)
            y = y + np.random.default_rng(700 + i).normal(0, NOISE, GRID.size)
            ids.append(self.p.add_spectrum(Spectrum(GRID, y, name=f"Tablet 20:1 + 5 B ({i + 1})",
                                                    concentrations={"A": 20.0, "B": 6.0}),
                                           tid, role="sample"))
        out = self.univariate("enrich", "B", "Constant multiplication / SS-CM (extended Y)",
                              refs={0: {"divisor": "B′"}},
                              params={0: {"start": 310.0, "end": 350.0}}, meas={"w1": 268.0},
                              mixtures=ids, spike=5.0,
                              notes=["Sample A 20 + B 1 µg/mL (ratio 20 : 1, B below its range) "
                                     "spiked with 5 µg/mL B; 'Found − added' is the B content "
                                     "(true 1 µg/mL)."])
        return out


# --------------------------------------------------------------------------- #
# Sample Excel workbooks
# --------------------------------------------------------------------------- #
SHEET_LABEL = {"Standards": "Std", "Divisor": "Divisor", "Mixtures": "Mix", "Samples": "Sample",
               "Calibration": "Train"}


def workbook_name(guide: "Guide", rid: int, sheet: str, k: int) -> str:
    """Column title that 'Fill concentrations from names' understands."""
    rec = guide.p.record(rid)
    if rec.kind == "derived":
        if "unit concentration" in rec.name:
            c = next(c for c, v in rec.concentrations.items() if v)
            return f"Unit divisor: {c} 1"
        return "Second divisor: D1 of the Z/Y ratio"
    comps = [c for c in SYSTEMS[SYSTEM_OF[next(iter(rec.concentrations))]]["pure"]
             if c in rec.concentrations]
    conc = ", ".join(f"{c} {rec.concentrations[c]:g}" for c in comps)
    label = SHEET_LABEL.get(sheet, sheet)
    if sheet == "Calibration":
        label = f"Train {k}"
    return f"{label}: {conc}"


def write_workbook(path: Path, guide: "Guide", e: "Entry") -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    r, data = e.result, e.result.get("data") or {}
    bold, head = Font(bold=True), PatternFill("solid", fgColor="DCE6F2")
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "_Read me"
    names: dict[int, str] = {}
    sheets_written = []
    # spectra sheets (raw spectra by role; processed divisors on their own sheet)
    groups: list[tuple[str, list[int]]] = []
    for sheet, ids in data.get("sheets", {}).items():
        raw = [i for i in ids if guide.p.record(i).kind == "raw"]
        der = [i for i in ids if guide.p.record(i).kind == "derived"]
        if raw:
            groups.append((sheet, raw))
        for j, i in enumerate(der):
            groups.append((f"Divisor (processed{'' if j == 0 else f' {j + 1}'})", [i]))
    for sheet, ids in groups:
        ws = wb.create_sheet(sheet)
        spectra = guide.p.spectra(ids)
        ws.cell(1, 1, "Wavelength (nm)").font = bold
        for j, (rid, sp) in enumerate(zip(ids, spectra), start=2):
            names[rid] = workbook_name(guide, rid, sheet.split(" (")[0], j - 1)
            c = ws.cell(1, j, names[rid])
            c.font, c.fill = bold, head
            ws.column_dimensions[get_column_letter(j)].width = max(14, len(names[rid]) + 2)
        x = spectra[0].wavelengths
        for i, w in enumerate(x, start=2):
            ws.cell(i, 1, round(float(w), 2))
            for j, sp in enumerate(spectra, start=2):
                ws.cell(i, j, round(float(np.interp(w, sp.wavelengths, sp.values)), 6))
        ws.column_dimensions["A"].width = 16
        ws.freeze_panes = "B2"
        sheets_written.append((sheet, len(ids)))
    # calculation with Excel formulas (single-signal methods)
    calc = data.get("calc")
    if calc:
        ws = wb.create_sheet("_Calculation")
        ws["A1"] = "Calibration and determination with Excel formulas"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Signal = {calc['signal']} (measured by Spectro on the processed spectra)"
        ws["A3"] = ("Change nothing here: the formulas show how Spectro turns signals into "
                    "concentrations. Compare with the _Expected results sheet.")
        for j, h in enumerate(["Standard", f"{calc['compound']} (µg/mL)", "Signal"], 1):
            c = ws.cell(5, j, h)
            c.font, c.fill = bold, head
        r0 = 6
        for k, (rid, conc, sig) in enumerate(calc["cal"]):
            ws.cell(r0 + k, 1, names.get(rid, guide.p.record(rid).name))
            ws.cell(r0 + k, 2, conc)
            ws.cell(r0 + k, 3, round(float(sig), 8))
        r1 = r0 + len(calc["cal"]) - 1
        rs, ri, rr = r1 + 2, r1 + 3, r1 + 4
        for row, lab, f in ((rs, "Slope", f"=SLOPE(C{r0}:C{r1},B{r0}:B{r1})"),
                            (ri, "Intercept", f"=INTERCEPT(C{r0}:C{r1},B{r0}:B{r1})"),
                            (rr, "Correlation coefficient r", f"=CORREL(B{r0}:B{r1},C{r0}:C{r1})")):
            ws.cell(row, 1, lab).font = bold
            ws.cell(row, 2, f)
        t0 = rr + 3
        hdr = ["Mixture", f"{calc['compound']} taken (µg/mL)", "Signal",
               "Found = (Signal − Intercept) / Slope" if not calc["direct"]
               else "Found = Signal (concentration value)", "Recovery %"]
        if calc["spike"]:
            hdr.append(f"Found − added ({calc['spike']:g})")
        for j, h in enumerate(hdr, 1):
            c = ws.cell(t0, j, h)
            c.font, c.fill = bold, head
            c.alignment = Alignment(wrap_text=True)
        for k, (rid, taken, sig) in enumerate(calc["test"]):
            row = t0 + 1 + k
            ws.cell(row, 1, names.get(rid, guide.p.record(rid).name))
            ws.cell(row, 2, taken)
            ws.cell(row, 3, round(float(sig), 8))
            ws.cell(row, 4, f"=C{row}" if calc["direct"] else f"=(C{row}-$B${ri})/$B${rs}")
            ws.cell(row, 5, f'=IF(B{row}>0,D{row}/B{row}*100,"")')
            if calc["spike"]:
                ws.cell(row, 6, f"=D{row}-{calc['spike']}")
        t1 = t0 + len(calc["test"])
        ws.cell(t1 + 2, 4, "Mean recovery %").font = bold
        ws.cell(t1 + 2, 5, f"=AVERAGE(E{t0 + 1}:E{t1})")
        ws.cell(t1 + 3, 4, "RSD %").font = bold
        ws.cell(t1 + 3, 5, f"=STDEV(E{t0 + 1}:E{t1})/AVERAGE(E{t0 + 1}:E{t1})*100")
        for col, w in zip("ABCDEF", (30, 18, 14, 26, 14, 16)):
            ws.column_dimensions[col].width = w
    # expected results
    ws = wb.create_sheet("_Expected results")
    ws["A1"] = f"What Spectro gives for this example — {e.title}"
    ws["A1"].font = Font(bold=True, size=13)
    row = 3
    for line in r.get("settings", []):
        ws.cell(row, 1, line.replace("*", ""))
        row += 1
    row += 1
    for j, h in enumerate(r.get("headers", []), 1):
        c = ws.cell(row, j, h)
        c.font, c.fill = bold, head
    for rr_ in r.get("rows", []):
        row += 1
        for j, v in enumerate(rr_, 1):
            try:
                ws.cell(row, j, float(v))
            except (TypeError, ValueError):
                ws.cell(row, j, v)
    row += 2
    for line in r.get("summary", []):
        ws.cell(row, 1, line).font = bold
        row += 1
    ws.column_dimensions["A"].width = 34
    # read me
    b = BEGINNER.get(e.key, {})
    comps = sorted({c for g in groups for i in g[1]
                    for c in guide.p.record(i).concentrations},
                   key=lambda c: "ABXYZUVW".index(c))
    lines = [
        (f"{e.title}", Font(bold=True, size=14)),
        (f"Family: {CATEGORIES[e.category]}", Font(italic=True)),
        ("", None),
        ("The idea in plain words", bold), (b.get("plain", e.principle), None),
        ("", None),
        ("What is in this workbook", bold)]
    lines += [(f"• Sheet '{n}': {k} spectr{'um' if k == 1 else 'a'} (wavelength in column A, "
               "one spectrum per column; the column title holds the concentrations)", None)
              for n, k in sheets_written]
    if calc:
        lines.append(("• Sheet '_Calculation': the calibration and the results done with Excel "
                      "formulas, so every number can be followed.", None))
    lines += [("• Sheet '_Expected results': what Spectro gives for this example.", None),
              ("• Sheets whose name starts with _ are notes; Spectro does not import them.",
               None)]
    if data.get("note"):
        lines.append(("• " + data["note"], None))
    lines += [("", None), ("How to use it in Spectro", bold),
              ("1. File → New project…", None),
              (f"2. Edit → Compounds… and add: {', '.join(comps)}", None),
              ("3. File → Import spectra… → Add files… → choose this workbook.", None),
              ("4. Press 'Fill concentrations from names'. The Role column is filled from the "
               "sheet names (Standards → standard, Mixtures → mixture, …). Press Import.", None)]
    lines += [(f"{k}. {st.replace('*', '')}", None) for k, st in enumerate(e.steps, 5)]
    lines += [("   (In these steps X is the compound being determined and Y, Z the interfering "
               f"ones — in this workbook the compounds are {', '.join(comps)}.)", None)]
    lines += [(f"{len(e.steps) + 5}. Compare your results with the sheet '_Expected results'.",
               None),
              ("", None), ("Settings used in this example", bold)]
    lines += [("• " + x.replace("*", ""), None) for x in r.get("settings", [])]
    lines += [("", None), ("Checks / common mistakes", bold), (b.get("check", ""), None),
              ("", None),
              ("The spectra are simulated (Gaussian bands + 0.0004 AU noise) to show the method "
               "clearly; they are not measured data. See docs/methods/README.md for the full "
               "guide with screenshots.", Font(italic=True, color="666666"))]
    for i, (text, font) in enumerate(lines, 1):
        c = info.cell(i, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font
    info.column_dimensions["A"].width = 120
    wb.save(path)


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
def md_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return out


def _images(r: dict) -> list[tuple[str, str]]:
    if r.get("images"):
        return r["images"]
    return [(r["image"], "")] if r.get("image") else []


INTRO = ("How to apply every spectrophotometric method in Spectro — written for "
         "beginners. Each method has: the idea in plain words, the principle, what you need "
         "to measure, step-by-step instructions, checks and common mistakes, a worked example "
         "with a 'why it works' figure and screenshots of every step, and a **sample Excel "
         "workbook** with the example's spectra (import it and reproduce the result). The "
         "examples use simulated systems (Gaussian bands, 0.0004 AU noise) — not measured "
         "data — and every number was produced by the app itself when this guide was "
         "generated (`python tools/method_guide.py`).")


def write_markdown(out: Path, entries: list[Entry], primer_images=()) -> None:
    md = ["# Spectro — method guide", "", INTRO, "",
          "**Tip:** *Tools → Method optimizer* tries most of these methods automatically on your "
          "standards and ranks them; this guide explains what each one does and how to set it "
          "up by hand.", "", "## Getting started", "", PRIMER.strip(), ""]
    for img, cap in primer_images:
        md += [f"![{cap}]({img})", "", f"*{cap}*", ""]
    md += ["### 6. The guide systems", ""]
    for name, s in SYSTEMS.items():
        md.append(f"* **{name}** — {s['about']} Standards "
                  f"{s['levels'][0]}–{s['levels'][-1]} µg/mL; divisors "
                  + ", ".join(f"{c}′ = {v} µg/mL" for c, v in s["divisors"].items()) + ".")
    md += ["", "![Pure spectra of the guide systems](img/systems.png)", "", "## Contents", ""]
    for cat, title in CATEGORIES.items():
        items = [e for e in entries if e.category == cat]
        if items:
            md.append(f"* **{title}** — " + ", ".join(
                f"[{e.title}](#{e.key})" for e in items))
    for cat, title in CATEGORIES.items():
        items = [e for e in entries if e.category == cat]
        if not items:
            continue
        md += ["", f"## {title}", "", DESCRIPTIONS[cat], ""]
        for e in items:
            b = BEGINNER.get(e.key, {})
            md += [f'<a id="{e.key}"></a>', "", f"### {e.title}", ""]
            if b.get("plain"):
                md += [f"**In plain words.** {b['plain']}", ""]
            md += [f"**Principle.** {e.principle}", "", f"**Use when.** {e.when}", ""]
            if b.get("need"):
                md += [f"**What you need.** {b['need']}", ""]
            md += ["**How to apply in Spectro**", ""]
            md += [f"{i}. {st}" for i, st in enumerate(e.steps, 1)]
            if b.get("check"):
                md += ["", f"**Checks and common mistakes.** {b['check']}"]
            r = e.result
            if r.get("error"):
                md += ["", f"*Example could not be generated: {r['error']}*"]
            elif r:
                md += ["", "**Worked example**", ""]
                if (out / "data" / f"{e.key}.xlsx").exists():
                    md += [f"📊 **Sample data:** [data/{e.key}.xlsx](data/{e.key}.xlsx) — "
                           "import it (see *Getting started → 4*) and reproduce the numbers "
                           "below.", ""]
                md += [f"* {x}" for x in r.get("settings", [])]
                imgs = _images(r)
                if imgs and imgs[0][0].endswith("_concept.png"):
                    md += ["", f"![{e.title} — why it works]({imgs[0][0]})", "",
                           f"*{imgs[0][1]}*"]
                    imgs = imgs[1:]
                if r.get("rows"):
                    md += [""] + md_table(r["headers"], r["rows"])
                if r.get("summary"):
                    md += ["", "Result: " + "; ".join(r["summary"]) + "."]
                for img, cap in imgs:
                    md += ["", f"![{e.title}]({img})", ""] + ([f"*{cap}*"] if cap else [])
            if e.refs:
                md += ["", "*Literature:* " + "; ".join(REFS.get(k.strip(), k.strip())
                                                        for k in e.refs.split(","))]
    (out / "README.md").write_text("\n".join(md) + "\n", encoding="utf-8")


def _inline(t: str) -> str:
    import re
    t = html.escape(t)
    t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    return re.sub(r"\*(.+?)\*", r"<i>\1</i>", t)


def md_to_html(text: str) -> str:
    """Enough Markdown for the primer: ### headings, lists, tables, paragraphs."""
    out, para, lst, table = [], [], None, []

    def flush():
        nonlocal para, lst, table
        if para:
            out.append("<p>" + _inline(" ".join(para)) + "</p>")
        if lst:
            tag, items = lst
            out.append(f"<{tag}>" + "".join(f"<li>{_inline(i)}</li>" for i in items)
                       + f"</{tag}>")
        if table:
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in table
                    if not set(r.replace("|", "").strip()) <= {"-"}]
            out.append("<table>" + "".join(
                "<tr>" + "".join(f"<{'th' if k == 0 else 'td'}>{_inline(c)}"
                                 f"</{'th' if k == 0 else 'td'}>" for c in row) + "</tr>"
                for k, row in enumerate(rows)) + "</table>")
        para, lst, table = [], None, []
    import re
    for line in text.splitlines():
        st = line.strip()
        if not st:
            flush()
        elif st.startswith("### "):
            flush()
            out.append(f"<h3>{_inline(st[4:])}</h3>")
        elif st.startswith("|"):
            if para or lst:
                flush()
            table.append(st)
        elif st.startswith("* ") or re.match(r"\d+\. ", st):
            tag = "ul" if st.startswith("* ") else "ol"
            item = st[2:] if tag == "ul" else st.split(". ", 1)[1]
            if para or table:
                flush()
            if lst and lst[0] == tag:
                lst[1].append(item)
            else:
                flush()
                lst = (tag, [item])
        elif lst and line.startswith("  "):
            lst[1][-1] += " " + st
        else:
            if lst or table:
                flush()
            para.append(st)
    flush()
    return "".join(out)


def write_html(out: Path, entries: list[Entry], primer_images=()) -> None:
    e_ = html.escape
    parts = ["<!doctype html><html><head><meta charset='utf-8'>"
             "<meta name='viewport' content='width=device-width,initial-scale=1'>"
             "<title>Spectro method guide</title><style>"
             ":root{--bg:#fcfcfb;--fg:#0b0b0b;--mut:#52514e;--line:#e3e2de;--card:#fff;"
             "--acc:#2a78d6;--soft:#eef4fc}"
             "@media (prefers-color-scheme:dark){:root{--bg:#161615;--fg:#eceae6;--mut:#b5b3ad;"
             "--line:#34332f;--card:#1f1f1d;--acc:#6aa5ee;--soft:#1d2a3a}}"
             "body{font-family:Segoe UI,Arial,sans-serif;margin:0;background:var(--bg);"
             "color:var(--fg);line-height:1.55}main{max-width:1100px;margin:auto;padding:24px}"
             "nav{background:var(--card);border:1px solid var(--line);border-radius:8px;"
             "padding:12px 18px}h2{border-bottom:2px solid var(--acc);padding-bottom:4px;"
             "margin-top:48px}section{background:var(--card);border:1px solid var(--line);"
             "border-radius:8px;padding:6px 20px 16px;margin:20px 0}.plain{background:var(--soft);"
             "border-left:4px solid var(--acc);padding:8px 12px;border-radius:4px}"
             "figure{margin:14px 0}figcaption{color:var(--mut);font-size:14px}"
             "img{max-width:100%;border:1px solid var(--line);border-radius:6px}"
             "table{border-collapse:collapse;font-size:13px;margin:8px 0;display:block;"
             "overflow-x:auto}td,th{border:1px solid var(--line);padding:3px 8px;text-align:left}"
             ".mut{color:var(--mut)}a{color:var(--acc)}.data{font-weight:600}"
             "</style></head><body><main><h1>Spectro — method guide</h1>"
             f"<p class='mut'>{_inline(INTRO)}</p><h2>Getting started</h2>",
             md_to_html(PRIMER)]
    for img, cap in primer_images:
        parts.append(f"<figure><img loading='lazy' src='{img}' alt='{e_(cap)}'>"
                     f"<figcaption>{_inline(cap)}</figcaption></figure>")
    parts.append("<h3>6. The guide systems</h3><ul>")
    for name, s in SYSTEMS.items():
        parts.append(f"<li><b>{e_(name)}</b> — {e_(s['about'])}</li>")
    parts.append("</ul><img src='img/systems.png' alt='Pure spectra of the guide systems'><nav>")
    for cat, title in CATEGORIES.items():
        items = [e for e in entries if e.category == cat]
        if items:
            parts.append(f"<p><b>{e_(title)}</b><br>" + " · ".join(
                f"<a href='#{e.key}'>{e_(e.title)}</a>" for e in items) + "</p>")
    parts.append("</nav>")
    for cat, title in CATEGORIES.items():
        items = [e for e in entries if e.category == cat]
        if not items:
            continue
        parts.append(f"<h2>{e_(title)}</h2><p class='mut'>{e_(DESCRIPTIONS[cat])}</p>")
        for e in items:
            r, b = e.result, BEGINNER.get(e.key, {})
            parts.append(f"<section id='{e.key}'><h3>{e_(e.title)}</h3>")
            if b.get("plain"):
                parts.append(f"<p class='plain'><b>In plain words.</b> {_inline(b['plain'])}</p>")
            parts.append(f"<p><b>Principle.</b> {_inline(e.principle)}</p>"
                         f"<p><b>Use when.</b> {_inline(e.when)}</p>")
            if b.get("need"):
                parts.append(f"<p><b>What you need.</b> {_inline(b['need'])}</p>")
            parts.append("<b>How to apply</b><ol>" + "".join(
                f"<li>{_inline(st)}</li>" for st in e.steps) + "</ol>")
            if b.get("check"):
                parts.append(f"<p><b>Checks and common mistakes.</b> {_inline(b['check'])}</p>")
            if r.get("error"):
                parts.append(f"<p><i>Example could not be generated: {e_(r['error'])}</i></p>")
            elif r:
                parts.append("<h4>Worked example</h4>")
                if (out / "data" / f"{e.key}.xlsx").exists():
                    parts.append(f"<p class='data'>📊 Sample data: <a href='data/{e.key}.xlsx'>"
                                 f"{e.key}.xlsx</a> — import it and reproduce the numbers.</p>")
                parts.append("<ul>" + "".join(f"<li>{_inline(x)}</li>"
                                              for x in r.get("settings", [])) + "</ul>")
                imgs = _images(r)
                if imgs and imgs[0][0].endswith("_concept.png"):
                    parts.append(f"<figure><img loading='lazy' src='{imgs[0][0]}' alt='concept'>"
                                 f"<figcaption>{_inline(imgs[0][1])}</figcaption></figure>")
                    imgs = imgs[1:]
                if r.get("rows"):
                    parts.append("<table><tr>" + "".join(f"<th>{e_(h)}</th>"
                                                         for h in r["headers"]) + "</tr>"
                                 + "".join("<tr>" + "".join(f"<td>{e_(str(c))}</td>"
                                                            for c in row) + "</tr>"
                                           for row in r["rows"]) + "</table>")
                if r.get("summary"):
                    parts.append(f"<p>Result: {e_('; '.join(r['summary']))}.</p>")
                for img, cap in imgs:
                    parts.append(f"<figure><img loading='lazy' src='{img}' alt='{e_(e.title)}'>"
                                 f"<figcaption>{_inline(cap)}</figcaption></figure>")
            if e.refs:
                parts.append("<p class='mut'>Literature: " + e_("; ".join(
                    REFS.get(k.strip(), k.strip()) for k in e.refs.split(","))) + "</p>")
            parts.append("</section>")
    parts.append("</main></body></html>")
    (out / "index.html").write_text("".join(parts), encoding="utf-8")


def systems_figure(out: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, axes = plt.subplots(1, 3, figsize=(13, 3.6), sharey=False)
    colours = ["#2a78d6", "#d6862a", "#3a9a5b"]
    for ax, (name, s) in zip(axes, SYSTEMS.items()):
        for col, (c, u) in zip(colours, s["pure"].items()):
            ax.plot(GRID, u * 10, color=col, lw=1.8, label=f"{c} 10 µg/mL")
        ax.set_title(name, fontsize=11)
        ax.set_xlabel("Wavelength (nm)")
        ax.set_xlim(200, 400)
        ax.legend(frameon=False, fontsize=9)
        ax.grid(alpha=0.2)
    axes[0].set_ylabel("Absorbance")
    fig.tight_layout()
    fig.savefig(out / "img" / "systems.png", dpi=110)
    plt.close(fig)
    compact_png(out / "img" / "systems.png")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--out", default=str(ROOT / "docs" / "methods"))
    ap.add_argument("--only", default="", help="comma-separated entry keys")
    args = ap.parse_args(argv)
    app = QApplication.instance() or QApplication(sys.argv[:1])
    app.setStyle("Fusion")
    out = Path(args.out)
    only = {k for k in args.only.split(",") if k}
    failures = []
    with tempfile.TemporaryDirectory() as tmp:
        guide = Guide(out, Path(tmp), app)
        systems_figure(out)
        for e in ENTRIES:
            if only and e.key not in only:
                continue
            print(f"{e.key}…", flush=True)
            try:
                e.result = getattr(guide, e.example)()
                for s in e.result.get("summary", []):
                    print("   ", s)
            except Exception as exc:
                traceback.print_exc()
                e.result = {"error": str(exc)}
                failures.append(e.key)
        for e in ENTRIES:
            if e.result.get("data"):
                try:
                    write_workbook(out / "data" / f"{e.key}.xlsx", guide, e)
                except Exception:
                    traceback.print_exc()
                    failures.append(f"{e.key} (workbook)")
        guide.win.close_project()
        primer = []
        if (out / "data" / "rd.xlsx").exists():
            try:
                primer = guide.primer_shots(out / "data" / "rd.xlsx", Path(tmp))
            except Exception:
                traceback.print_exc()
                failures.append("primer screenshots")
    entries = [e for e in ENTRIES if not only or e.key in only]
    write_markdown(out, entries, primer)
    write_html(out, entries, primer)
    print(f"\n{len(entries)} methods → {out / 'README.md'} and index.html")
    if failures:
        print("FAILED:", ", ".join(failures))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
